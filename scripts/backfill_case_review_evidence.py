"""回填用例审查产出丢失的证据（findings / candidates）。

背景
----
`testcases/review_service.py` 的 `_write_report` 会整体覆盖 `review.summary`，
把 `_checkpoint`（问题明细）丢掉；而旧版发布逻辑正是从 summary 里取 issues，
结果写进飞轮的产出 `issues_count=0`、`trace.candidates=[]`——摘要里有 14 条问题，
飞轮里一条证据都没有，评测/归因/金标全部失去依据。

修复已让新流程直接传入真实 issues。本脚本负责把**历史已发布产出**原地补齐：
从各自生成的 Excel 报告「问题明细」页解析问题，重建 findings 并写回
`GenerationOutput.content` / `metadata.protocol` 与 `RetrievalTrace.candidates`，
保持 output id 与 trace id 不变（幂等，可重复执行）。
"""
from __future__ import annotations

import json
import os

import django

os.environ.setdefault("DJANGO_SETTINGS_MODULE", "wharttest_django.settings")
django.setup()

import openpyxl  # noqa: E402
from knowledge_evolution.models import GenerationOutput  # noqa: E402
from testcases.models import TestCaseReview  # noqa: E402
from testcases.review_service import build_review_findings  # noqa: E402

ISSUE_HEADERS = {
    "Sheet": "sheet",
    "行号": "row",
    "用例编号/名称": "case_id",
    "模块": "module",
    "原文": "original",
    "严重程度": "severity",
    "问题类型": "issue_type",
    "问题说明": "description",
    "修改建议": "suggestion",
    "判定": "judgement",
}


def read_report(path: str):
    """从 Excel 报告恢复 issues / pending / governance。"""
    workbook = openpyxl.load_workbook(path, data_only=True)
    issues: list[dict] = []
    sheet = workbook["问题明细"] if "问题明细" in workbook.sheetnames else None
    if sheet is not None:
        header = [str(c.value or "").strip() for c in sheet[1]]
        index = {name: pos for pos, name in enumerate(header)}
        for row in sheet.iter_rows(min_row=2, values_only=True):
            if not any(value not in (None, "") for value in row):
                continue
            item = {}
            for label, key in ISSUE_HEADERS.items():
                pos = index.get(label)
                item[key] = row[pos] if pos is not None and pos < len(row) else ""
            if item.get("row") in (None, ""):
                continue
            issues.append(item)
    pending = []
    if "待确认事项" in workbook.sheetnames:
        for row in workbook["待确认事项"].iter_rows(min_row=2, values_only=True):
            value = row[1] if len(row) > 1 else None
            if value:
                pending.append(str(value))
    governance = []
    if "治理建议" in workbook.sheetnames:
        for row in workbook["治理建议"].iter_rows(min_row=2, values_only=True):
            value = row[1] if len(row) > 1 else None
            if value:
                governance.append(str(value))
    return issues, pending, governance


def rebuild_content(review, issues, findings, pending, governance):
    """复刻发布时的 output 结构，保证正文与协议一致。

    与代码审查一致：正文携带结论明细本体（findings），而不只是计数，
    否则飞轮的 L1 裁判看不到「问题说明/修改建议」，评测恒判不合格。
    """
    summary = {
        key: value
        for key, value in (review.summary or {}).items()
        if key not in {"_checkpoint", "trace_id", "output_id"}
    }
    output = {
        "summary": summary,
        "findings": findings,
        "pending": [str(item) for item in pending],
        "governance": [str(item) for item in governance],
        "issues_count": len(issues),
        "pending_count": len(pending),
        "governance_count": len(governance),
    }
    return output


def backfill(review, output):
    report = review.report_file
    if not report or not os.path.exists(report.path):
        print(f"  [skip] review {review.pk}: 无报告文件")
        return False
    issues, pending, governance = read_report(report.path)
    if not issues:
        print(f"  [skip] review {review.pk}: 报告里没有问题明细")
        return False

    findings = build_review_findings(review, issues)
    output_dict = rebuild_content(review, issues, findings, pending, governance)
    content = json.dumps(output_dict, ensure_ascii=False, sort_keys=True)
    import hashlib
    output_hash = hashlib.sha256(content.encode("utf-8")).hexdigest()

    before = len((output.trace.candidates or []) if output.trace else [])

    output.content = content
    output.output_hash = output_hash
    metadata = dict(output.metadata or {})
    protocol = dict(metadata.get("protocol") or {})
    protocol["findings"] = findings
    protocol["metrics"] = {
        "latency_ms": (protocol.get("metrics") or {}).get("latency_ms", 0),
        "token_usage": (protocol.get("metrics") or {}).get("token_usage", 0),
        "chunks": (protocol.get("metrics") or {}).get("chunks", 0),
        "uncovered_chunks": (protocol.get("metrics") or {}).get("uncovered_chunks", 0),
        "issues_count": len(issues),
        "pending_count": len(pending),
    }
    protocol["output_descriptor"] = {
        "content_hash": hashlib.sha256(
            json.dumps(output_dict, ensure_ascii=False, sort_keys=True, separators=(",", ":")).encode()
        ).hexdigest(),
        "keys": sorted(output_dict.keys()),
        "size_bytes": len(content.encode()),
    }
    metadata["protocol"] = protocol
    output.metadata = metadata
    output.save(update_fields=["content", "output_hash", "metadata"])

    if output.trace:
        trace = output.trace
        trace.candidates = findings
        trace.save(update_fields=["candidates"])
    print(f"  [ok] review {review.pk}: findings {before} -> {len(findings)} (issues={len(issues)}, pending={len(pending)}, governance={len(governance)})")
    return True


def main():
    print("=== 回填用例审查产出的问题明细证据 ===")
    changed = 0
    for review in TestCaseReview.objects.filter(status="completed").order_by("id"):
        output_id = (review.summary or {}).get("output_id")
        output = None
        if output_id:
            output = GenerationOutput.objects.filter(pk=output_id).select_related("trace").first()
        if output is None:
            output = (
                GenerationOutput.objects.filter(task_type="case_review", task_id=str(review.pk))
                .select_related("trace")
                .order_by("-created_at")
                .first()
            )
        if output is None:
            print(f"  [skip] review {review.pk}: 飞轮里没有对应产出")
            continue
        try:
            parsed = json.loads(output.content or "{}")
        except Exception:
            parsed = {}
        if isinstance(parsed, dict) and parsed.get("findings"):
            print(f"  [skip] review {review.pk}: 正文已含 {len(parsed['findings'])} 条结论明细，无需回填")
            continue
        if backfill(review, output):
            changed += 1
    print(f"=== 完成，回填 {changed} 条产出 ===")


if __name__ == "__main__":
    main()
