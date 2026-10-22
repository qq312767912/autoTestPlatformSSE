from io import BytesIO
from types import SimpleNamespace
from unittest.mock import patch

from django.core.exceptions import ValidationError
from django.test import SimpleTestCase
from openpyxl import Workbook, load_workbook

from testcases.skill_report import declared_sheets, render_sheets, validate_sheets
from testcases.review_service import _review_chunk, _write_report


class SkillReportContractTests(SimpleTestCase):
    def declaration(self, name='独立核查'):
        return validate_sheets([{'name': name, 'columns': [{'key':'case', 'title':'用例'}, {'key':'status', 'title':'状态'}],
            'highlight':[{'column':'status','equals':'待确认','fill':'FFF2CC','color':'9C6500'}]}])

    def test_skill_sheets_merge_batches_highlight_and_preserve_text(self):
        book=Workbook()
        spec=self.declaration()
        row={'case':'=HYPERLINK("bad")', 'status':'待确认'}
        result=render_sheets(book,spec,[[{'name':'独立核查','rows':[row]}],[{'name':'独立核查','rows':[row]}]])
        self.assertEqual(result['extra_sheets'],['独立核查'])
        sheet=book['独立核查']
        self.assertEqual(sheet.max_row,4)
        self.assertEqual(sheet['B4'].fill.fgColor.rgb,'00FFF2CC')
        self.assertEqual(sheet['A4'].data_type,'s')
        self.assertEqual(sheet['A4'].value,row['case'])

    def test_missing_result_keeps_required_sheet_and_warning(self):
        book=Workbook()
        result=render_sheets(book,self.declaration(),[])
        self.assertIn('未返回',book['独立核查']['A4'].value)
        self.assertTrue(result['report_warnings'])
        self.assertEqual(render_sheets(Workbook(),[],[])['extra_sheets'],[])

    def test_dynamic_model_definition_and_invalid_structure(self):
        book=Workbook()
        result=render_sheets(book,[],[[{'name':'Skill文字要求的表','columns':['结论'],'rows':[{'结论':'有待确认项'}]}]])
        self.assertEqual(result['extra_sheets'],['Skill文字要求的表'])
        self.assertEqual(book['Skill文字要求的表']['A4'].value,'有待确认项')
        with self.assertRaises(ValidationError): validate_sheets([{'name':'非法/表','columns':['a']}])
        with self.assertRaises(ValidationError): declared_sheets('---\nname: review\nreport: bad\n---\n')

    def test_request_and_export_follow_this_skill(self):
        skill='---\nname: review\nreport:\n  sheets:\n    - name: 此包独有表\n      columns: [结论]\n---\n按输入填写结论。'
        with patch('testcases.review_service.safe_llm_invoke',return_value=SimpleNamespace(content='{}')) as invoke:
            _review_chunk(object(),skill,[],'')
            self.assertIn('此包独有表',invoke.call_args[0][1][1].content)
        stored={}
        def save(filename,content,save=False): stored['data']=content.read()
        review=SimpleNamespace(source_name='case.xlsx',skill_snapshot=skill,skill_name='review',business_context='',custom_rules='',get_review_mode_display=lambda:'标准',report_file=SimpleNamespace(save=save))
        meta=_write_report(review,[],[],[],[],extra_reports=[[{'name':'此包独有表','rows':[{'结论':'待确认'}]}]])
        book=load_workbook(BytesIO(stored['data']))
        self.assertIn('此包独有表',book.sheetnames)
        self.assertEqual(book['此包独有表']['A4'].value,'待确认')
        self.assertEqual(meta['extra_sheets'],['此包独有表'])
