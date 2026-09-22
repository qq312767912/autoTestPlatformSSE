"""基于 test-case-clarity-review Skill 的测试用例文件审查。"""

import csv
import hashlib
import json
import logging
import re
import threading
from collections import Counter
from concurrent.futures import ThreadPoolExecutor, as_completed
from datetime import datetime
from pathlib import Path
from time import monotonic, sleep

from django.core.files.base import ContentFile
from django.utils import timezone
from langchain_core.messages import HumanMessage, SystemMessage
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from requirements.services import create_llm_instance, safe_llm_invoke

from .models import TestCaseReview, TestCaseReviewLLMConfig

logger = logging.getLogger(__name__)

FALLBACK_SKILL = """审查测试用例的可执行性与验收清晰度。不要凭空补造业务规则；未知标准标记为待业务确认。按风险检查前置条件、测试数据、步骤、预期、业务结果、边界异常、数据一致性、证据与可维护性。"""
TESTCASE_REVIEW_CHUNK_SIZE = 20
TESTCASE_REVIEW_CHUNK_ATTEMPTS = 3
TESTCASE_REVIEW_MAX_ATTEMPTS = 5
TESTCASE_REVIEW_MAX_WORKERS = 2
TESTCASE_REVIEW_TOTAL_BUDGET_SECONDS = 45 * 60
TESTCASE_REVIEW_CIRCUIT_BREAKER = 3
TESTCASE_REVIEW_RETRY_BASE_DELAY = 2
TESTCASE_REVIEW_RETRY_MAX_DELAY = 30


def _get_testcase_review_llm_config():
    """只解析用例审查专属配置，禁止隐式回退到平台通用 LLM。

    用例审查要求模型按提示词稳定输出 JSON，与对话/编排对模型的偏好不同。
    若这里回退到通用配置，改通用配置会悄悄改变审查行为，且管理员无法为审查
    单独选一个"吐 JSON 更稳"的模型，因此宁可显式失败。
    """
    config = TestCaseReviewLLMConfig.objects.filter(is_active=True).first()
    if not config:
        raise ValueError("尚未配置并启用用例审查专用 LLM，请联系平台管理员配置")
    if not config.api_key:
        raise ValueError("用例审查专用 LLM 缺少 API Key，请联系平台管理员补全")
    return config


def _chunk_attempt_limit(config):
    """单批最大尝试次数：尊重专用配置的 max_retries，但设有下限与上限。

    内网模型网关会出现瞬时 5xx（如 502 upstream_unavailable）。历史实现每批
    只尝试 2 次且间隔固定 2 秒，一次上游抖动就让整项任务失败，故这里设下限；
    同时设上限，避免单批长时间占用模型。
    """
    try:
        configured = int(getattr(config, "max_retries", 0) or 0) + 1
    except (TypeError, ValueError):
        configured = TESTCASE_REVIEW_CHUNK_ATTEMPTS
    return max(TESTCASE_REVIEW_CHUNK_ATTEMPTS, min(configured, TESTCASE_REVIEW_MAX_ATTEMPTS))


def _retry_delay(attempt):
    """指数退避：2、4、8……并封顶，避免在上游故障窗口内高频重试。"""
    return min(TESTCASE_REVIEW_RETRY_MAX_DELAY, TESTCASE_REVIEW_RETRY_BASE_DELAY * (2 ** (attempt - 1)))


HEADER_ALIASES = {
    "identity": ("编号", "用例id", "用例编号", "用例名称", "测试用例", "caseid", "casename"),
    "module": ("模块", "所属用例库", "一级模块", "功能模块"),
    "precondition": ("前置条件", "前置", "测试数据"),
    "steps": ("步骤", "步骤描述", "操作步骤", "测试步骤"),
    "expected": ("预期", "预期结果", "期望结果", "期望输出", "验收标准"),
    "priority": ("优先级", "用例等级"),
}


def _normalized_header(value):
    return re.sub(r"[\s_\-/:：]+", "", str(value or "")).lower()


def _header_categories(cells):
    categories = set()
    for cell in cells:
        normalized = _normalized_header(cell)
        if not normalized:
            continue
        for category, aliases in HEADER_ALIASES.items():
            if any(_normalized_header(alias) == normalized for alias in aliases):
                categories.add(category)
                break
    return categories


def _is_case_header(cells):
    categories = _header_categories(cells)
    return (
        {"steps", "expected"}.issubset(categories)
        or ("identity" in categories and bool(categories & {"steps", "expected", "precondition"}))
    )


def _case_rows(sheet_name, source_rows):
    """在一个工作表中识别一个或多个用例表头，过滤封面、评审记录和分组标题行。"""
    rows = []
    headers = None
    for row_number, cells in source_rows:
        if _is_case_header(cells):
            headers = cells
            continue
        if headers is None:
            continue
        width = max(len(headers), len(cells))
        padded_headers = [*headers, *([""] * (width - len(headers)))]
        padded_cells = [*cells, *([""] * (width - len(cells)))]
        columns = {
            header: value
            for header, value in zip(padded_headers, padded_cells)
            if header and value
        }
        # 只有一个值的行通常是“Chrome 系列”等分组标题，不是用例。
        if len(columns) < 2:
            continue
        rows.append({
            "sheet": sheet_name,
            "row": row_number,
            "cells": cells,
            "columns": columns,
        })
    return rows


def build_skill_snapshot(skill=None):
    """固化本次审查所用 Skill，避免执行期间 Skill 被编辑或删除。"""
    if skill and skill.skill_content:
        content = skill.skill_content
        try:
            skill_path = skill.get_full_path()
            if not skill_path:
                return content
            references = Path(skill_path) / "references"
            if references.exists():
                for path in sorted(references.rglob("*.md")):
                    content += f"\n\n# 参考规则：{path.name}\n" + path.read_text(encoding="utf-8")
        except Exception:
            logger.exception("读取 Skill 参考规则失败: %s", skill.name)
        return content
    bundled = Path("/app/bundled_skills/test-case-clarity-review")
    skill_md = bundled / "SKILL.md"
    rules_md = bundled / "references" / "review-rules.md"
    if skill_md.exists():
        return skill_md.read_text(encoding="utf-8") + ("\n\n" + rules_md.read_text(encoding="utf-8") if rules_md.exists() else "")
    return FALLBACK_SKILL


def _skill_prompt(review):
    return review.skill_snapshot or build_skill_snapshot(review.selected_skill)


def _read_rows(path):
    suffix = Path(path).suffix.lower()
    rows = []
    if suffix == ".xlsx":
        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            for sheet in workbook.worksheets:
                source_rows = []
                for row_number, values in enumerate(sheet.iter_rows(values_only=True), 1):
                    cells = [str(value).strip() if value is not None else "" for value in values]
                    if any(cells):
                        source_rows.append((row_number, cells))
                rows.extend(_case_rows(sheet.title, source_rows))
        finally:
            workbook.close()
    else:
        with open(path, "r", encoding="utf-8-sig", errors="replace", newline="") as handle:
            source_rows = []
            for row_number, values in enumerate(csv.reader(handle), 1):
                cells = [str(value).strip() for value in values]
                if any(cells):
                    source_rows.append((row_number, cells))
            rows.extend(_case_rows("CSV", source_rows))
    return rows


def _extract_json(text):
    cleaned = re.sub(r"^```(?:json)?\s*|\s*```$", "", (text or "").strip(), flags=re.I | re.S)
    try:
        return json.loads(cleaned)
    except json.JSONDecodeError:
        match = re.search(r"\{.*\}", cleaned, flags=re.S)
        if not match:
            raise ValueError("模型未返回可解析的 JSON")
        return json.loads(match.group(0))


def _review_chunk(llm, skill_prompt, rows, business_context):
    schema = {
        "issues": [{
            "sheet": "Sheet", "row": 2, "case_id": "编号或名称", "module": "模块",
            "original": "原文", "severity": "高|中|低", "issue_type": "问题类型",
            "description": "问题说明", "suggestion": "具体修改建议", "judgement": "确定缺陷|待业务确认|优化建议",
            "rewrite": "可直接替换的改写示例或空字符串",
        }],
        "pending_confirmations": ["待确认事项"], "governance_suggestions": ["治理建议"],
    }
    prompt = (
        "请严格按下方 Skill 审查这些测试用例行。只输出一个 JSON 对象，不要 Markdown。"
        "同一根因在同一用例内合并；高风险必须语义复核；识别并跳过表头，不要把表头当作测试用例。"
        "每行 columns 已按识别到的原表头映射，cells 保留原列顺序。\n"
        f"业务背景：{business_context or '未提供；未知业务规则必须标为待业务确认'}\n"
        f"输出结构示例：{json.dumps(schema, ensure_ascii=False)}\n"
        f"待审查数据：{json.dumps(rows, ensure_ascii=False)}"
    )
    # 此处只调用一次：重试次数与指数退避统一由 run_testcase_review 的外层循环管理。
    # 若两层同时重试，单批最坏耗时会放大成 (1 + max_retries) 倍超时时间，
    # 无法用总时间预算约束，反而更容易撞上 Celery 时限。
    response = safe_llm_invoke(
        llm,
        [SystemMessage(content=skill_prompt), HumanMessage(content=prompt)],
        max_retries=1,
        retry_delay=2,
    )
    return _extract_json(response.content)


def _write_report(review, rows, issues, pending, governance, uncovered=None, chunks=None):
    uncovered = uncovered or {}
    chunks = chunks or []
    workbook = Workbook()
    info = workbook.active
    info.title = "审查摘要"
    counts = Counter(str(item.get("severity") or "中") for item in issues)
    type_counts = Counter(str(item.get("issue_type") or "其他") for item in issues)
    module_counts = Counter(str(item.get("module") or "未标注模块") for item in issues)
    issue_rows = {(str(item.get("sheet")), int(item.get("row") or 0)) for item in issues}
    estimated_cases = len(rows)
    uncovered_rows = sum(len(chunks[index - 1]) for index in uncovered if 0 < index <= len(chunks))
    top_issue = type_counts.most_common(1)[0][0] if type_counts else "未发现明确问题"
    summary_rows = [
        ("源文件", review.source_name), ("扫描时间", timezone.localtime().strftime("%Y-%m-%d %H:%M:%S")),
        ("用例数量（按表头后非空行估算）", estimated_cases), ("扫描非空行数", len(rows)),
        ("存在问题的用例/行", len(issue_rows)), ("问题总数", len(issues)),
        ("高风险", counts.get("高", 0)), ("中风险", counts.get("中", 0)), ("低风险", counts.get("低", 0)),
        ("最常见问题", top_issue),
        ("审查方式", review.get_review_mode_display()), ("使用 Skill", review.skill_name),
        ("业务背景", review.business_context or "未提供"), ("自定义审查规则", review.custom_rules or "未提供"),
        ("审查口径", f"{review.skill_name} + 用户自定义审查规则"),
        ("未覆盖批次", f"{len(uncovered)}/{len(chunks)}" if chunks else str(len(uncovered))),
        ("未覆盖用例行数", uncovered_rows),
    ]
    info.merge_cells("A1:B1")
    info["A1"] = "测试用例质量审查报告"
    info.merge_cells("A2:B2")
    info["A2"] = "审查结果仅用于提升用例的可执行性与验收清晰度，未知业务规则需由业务方确认。"
    info.append(["基本信息", "内容"])
    for key, value in summary_rows:
        info.append([key, value])
    info.column_dimensions["A"].width = 24
    info.column_dimensions["B"].width = 90

    distribution = workbook.create_sheet("问题分布")
    distribution.append(["维度", "分类", "数量"])
    for label, counter in (("严重程度", counts), ("问题类型", type_counts), ("模块", module_counts)):
        for name, count in counter.most_common():
            distribution.append([label, name, count])

    detail = workbook.create_sheet("问题明细")
    headers = [
        "Sheet", "行号", "用例编号/名称", "模块", "原文", "严重程度", "问题类型", "问题说明",
        "修改建议", "判定", "问题确认", "问题描述", "修改点", "不采纳原因",
    ]
    detail.append(headers)
    for item in issues:
        detail.append(
            [item.get(k, "") for k in [
                "sheet", "row", "case_id", "module", "original", "severity", "issue_type",
                "description", "suggestion", "judgement",
            ]] + ["", "", "", ""]
        )
    confirmation_validation = DataValidation(
        type="list", formula1='"是,否"', allow_blank=True,
        errorTitle="无效的问题确认值", error='请从下拉列表中选择“是”或“否”。',
        promptTitle="问题确认", prompt='确认采纳请选“是”，不采纳请选“否”。',
    )
    confirmation_validation.showErrorMessage = True
    confirmation_validation.showInputMessage = True
    if issues:
        detail.add_data_validation(confirmation_validation)
        confirmation_validation.add(f"K2:K{len(issues) + 1}")

    examples = workbook.create_sheet("改写示例")
    examples.append(["用例编号/名称", "原文", "改写示例"])
    for item in issues:
        if item.get("rewrite") and item.get("severity") in {"高", "中"}:
            examples.append([item.get("case_id", ""), item.get("original", ""), item.get("rewrite", "")])

    confirms = workbook.create_sheet("待确认事项")
    confirms.append(["序号", "待确认内容"])
    for index, value in enumerate(dict.fromkeys(pending), 1):
        confirms.append([index, value])
    governance_sheet = workbook.create_sheet("治理建议")
    governance_sheet.append(["序号", "建议"])
    for index, value in enumerate(dict.fromkeys(governance), 1):
        governance_sheet.append([index, value])

    confirmation_sheet = workbook.create_sheet("测试确认处理结果")
    confirmation_sheet.merge_cells("A1:F1")
    confirmation_sheet["A1"] = "测试确认处理结果"
    confirmation_sheet.merge_cells("A2:F2")
    confirmation_sheet["A2"] = "本页数据根据“问题明细”中的“问题确认”自动统计：“是”为采纳，“否”为不采纳。"
    confirmation_sheet.append([])
    confirmation_sheet.append(["总体统计", "数值"])
    issue_end_row = max(2, len(issues) + 1)
    confirmation_sheet.append(["共审查用例数", estimated_cases])
    confirmation_sheet.append(["审查提供意见", f"=COUNTA('问题明细'!$A$2:$A${issue_end_row})"])
    confirmation_sheet.append(["采纳", f'=COUNTIF(\'问题明细\'!$K$2:$K${issue_end_row},"是")'])
    confirmation_sheet.append(["不采纳", f'=COUNTIF(\'问题明细\'!$K$2:$K${issue_end_row},"否")'])
    confirmation_sheet.append(["待确认", "=B6-B7-B8"])
    confirmation_sheet.append(["采纳率", "=IF(B6=0,0,B7/B6)"])
    confirmation_sheet.append([])
    confirmation_sheet.append(["问题类型", "审查意见数", "采纳", "不采纳", "待确认", "采纳率"])
    for issue_type in type_counts:
        row_number = confirmation_sheet.max_row + 1
        confirmation_sheet.append([
            issue_type,
            f'=COUNTIF(\'问题明细\'!$G$2:$G${issue_end_row},A{row_number})',
            f'=COUNTIFS(\'问题明细\'!$G$2:$G${issue_end_row},A{row_number},\'问题明细\'!$K$2:$K${issue_end_row},"是")',
            f'=COUNTIFS(\'问题明细\'!$G$2:$G${issue_end_row},A{row_number},\'问题明细\'!$K$2:$K${issue_end_row},"否")',
            f"=B{row_number}-C{row_number}-D{row_number}",
            f"=IF(B{row_number}=0,0,C{row_number}/B{row_number})",
        ])

    navy = "1D3F66"
    blue = "2F75B5"
    pale_blue = "EAF2F8"
    stripe = "F7FAFC"
    border_color = "D6E0EA"
    thin_border = Border(
        left=Side(style="thin", color=border_color), right=Side(style="thin", color=border_color),
        top=Side(style="thin", color=border_color), bottom=Side(style="thin", color=border_color),
    )
    risk_styles = {
        "高": ("FCE8E6", "B42318"),
        "中": ("FFF2D6", "B54708"),
        "低": ("E8F3FF", "175CD3"),
    }

    for sheet in workbook.worksheets:
        sheet.sheet_view.showGridLines = False
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 0
        sheet.sheet_properties.tabColor = blue
        header_row = 3 if sheet is info else 1
        sheet.freeze_panes = f"A{header_row + 1}" if sheet.max_row > header_row else None
        if sheet is not info and sheet.max_row > 1:
            sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[header_row]:
            cell.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor=blue)
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = thin_border
        sheet.row_dimensions[header_row].height = 26
        for column in range(1, sheet.max_column + 1):
            sheet.column_dimensions[get_column_letter(column)].width = min(max(14, max(len(str(sheet.cell(row, column).value or "")) for row in range(1, min(sheet.max_row, 100) + 1)) + 2), 60)
        for row_index in range(header_row + 1, sheet.max_row + 1):
            fill = PatternFill("solid", fgColor=stripe if row_index % 2 == 0 else "FFFFFF")
            for cell in sheet[row_index]:
                cell.font = Font(name="Arial", size=10, color="26384A")
                cell.fill = fill
                cell.border = thin_border
                cell.alignment = Alignment(vertical="top", wrap_text=True)
            sheet.row_dimensions[row_index].height = 34 if sheet in {detail, examples} else 24

    info["A1"].fill = PatternFill("solid", fgColor=navy)
    info["A1"].font = Font(name="Arial", size=18, bold=True, color="FFFFFF")
    info["A1"].alignment = Alignment(vertical="center")
    info["B1"].fill = PatternFill("solid", fgColor=navy)
    info.row_dimensions[1].height = 44
    info["A2"].fill = PatternFill("solid", fgColor=pale_blue)
    info["B2"].fill = PatternFill("solid", fgColor=pale_blue)
    info["A2"].font = Font(name="Arial", size=10, italic=True, color="52677D")
    info["A2"].alignment = Alignment(vertical="center", wrap_text=True)
    info.row_dimensions[2].height = 34
    for row_index in range(4, info.max_row + 1):
        info.cell(row_index, 1).font = Font(name="Arial", size=10, bold=True, color="27445E")
        if info.cell(row_index, 1).value in risk_styles:
            fill_color, font_color = risk_styles[info.cell(row_index, 1).value]
            for cell in info[row_index]:
                cell.fill = PatternFill("solid", fgColor=fill_color)
            info.cell(row_index, 2).font = Font(name="Arial", size=11, bold=True, color=font_color)

    detail.page_setup.orientation = "landscape"
    detail.sheet_properties.tabColor = "C73E1D"
    detail.column_dimensions["A"].width = 16
    detail.column_dimensions["B"].width = 9
    detail.column_dimensions["C"].width = 28
    detail.column_dimensions["D"].width = 20
    detail.column_dimensions["E"].width = 48
    detail.column_dimensions["F"].width = 12
    detail.column_dimensions["G"].width = 24
    detail.column_dimensions["H"].width = 48
    detail.column_dimensions["I"].width = 48
    detail.column_dimensions["J"].width = 16
    detail.column_dimensions["K"].width = 14
    detail.column_dimensions["L"].width = 48
    detail.column_dimensions["M"].width = 36
    detail.column_dimensions["N"].width = 36
    for row_index in range(2, detail.max_row + 1):
        severity = str(detail.cell(row_index, 6).value or "")
        if severity in risk_styles:
            fill_color, font_color = risk_styles[severity]
            cell = detail.cell(row_index, 6)
            cell.fill = PatternFill("solid", fgColor=fill_color)
            cell.font = Font(name="Arial", size=10, bold=True, color=font_color)
            cell.alignment = Alignment(horizontal="center", vertical="center")
        judgement = str(detail.cell(row_index, 10).value or "")
        judgement_colors = {"确定缺陷": ("FCE8E6", "B42318"), "待业务确认": ("FFF2D6", "B54708"), "优化建议": ("EAF2F8", "175CD3")}
        if judgement in judgement_colors:
            fill_color, font_color = judgement_colors[judgement]
            cell = detail.cell(row_index, 10)
            cell.fill = PatternFill("solid", fgColor=fill_color)
            cell.font = Font(name="Arial", size=10, bold=True, color=font_color)
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    confirmation_sheet.sheet_properties.tabColor = "A64B4B"
    confirmation_sheet.freeze_panes = "A4"
    confirmation_sheet.auto_filter.ref = (
        f"A12:F{confirmation_sheet.max_row}" if confirmation_sheet.max_row > 12 else None
    )
    confirmation_sheet.column_dimensions["A"].width = 28
    for column in "BCDE":
        confirmation_sheet.column_dimensions[column].width = 16
    confirmation_sheet.column_dimensions["F"].width = 16
    confirmation_sheet["A1"].fill = PatternFill("solid", fgColor="8B3F3F")
    confirmation_sheet["A1"].font = Font(name="Arial", size=18, bold=True, color="FFFFFF")
    confirmation_sheet["A1"].alignment = Alignment(vertical="center")
    confirmation_sheet.row_dimensions[1].height = 40
    confirmation_sheet["A2"].fill = PatternFill("solid", fgColor="F4E9E7")
    confirmation_sheet["A2"].font = Font(name="Arial", size=10, italic=True, color="6B3530")
    confirmation_sheet["A2"].alignment = Alignment(vertical="center", wrap_text=True)
    confirmation_sheet.row_dimensions[2].height = 32
    for header_row in (4, 12):
        for cell in confirmation_sheet[header_row]:
            cell.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="8B3F3F")
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = thin_border
    for row_index in range(5, confirmation_sheet.max_row + 1):
        if row_index in {11, 12}:
            continue
        for cell in confirmation_sheet[row_index]:
            cell.border = thin_border
            cell.alignment = Alignment(vertical="center", wrap_text=True)
    confirmation_sheet["B10"].number_format = "0.00%"
    for row_index in range(13, confirmation_sheet.max_row + 1):
        confirmation_sheet.cell(row_index, 6).number_format = "0.00%"

    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True
    workbook.calculation.calcMode = "auto"

    for row_index in range(2, distribution.max_row + 1):
        if distribution.cell(row_index, 1).value == "严重程度":
            severity = str(distribution.cell(row_index, 2).value or "")
            if severity in risk_styles:
                fill_color, font_color = risk_styles[severity]
                for cell in distribution[row_index]:
                    cell.fill = PatternFill("solid", fgColor=fill_color)
                distribution.cell(row_index, 3).font = Font(name="Arial", size=10, bold=True, color=font_color)

    from io import BytesIO
    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    stem = Path(review.source_name).stem
    filename = f"{stem}_测试用例质量审查报告.xlsx"
    review.report_file.save(filename, ContentFile(output.read()), save=False)
    return {
        "total_rows": estimated_cases, "scanned_rows": len(rows), "problem_cases": len(issue_rows),
        "issue_count": len(issues), "high": counts.get("高", 0), "medium": counts.get("中", 0),
        "low": counts.get("低", 0),
        "uncovered_chunks": len(uncovered), "uncovered_rows": uncovered_rows,
        "total_chunks": len(chunks),
    }


def run_testcase_review(review_id):
    review = TestCaseReview.objects.get(pk=review_id)
    review.status = "running"
    review.started_at = timezone.now()
    review.current_step = "读取测试用例"
    review.progress = 5
    review.save(update_fields=["status", "started_at", "current_step", "progress", "updated_at"])
    rows = _read_rows(review.source_file.path)
    if not rows:
        raise ValueError("文件中没有可审查的非空内容")
    config = _get_testcase_review_llm_config()
    skill_prompt = _skill_prompt(review)
    if review.custom_rules.strip():
        skill_prompt += "\n\n# 本次用户指定的审查规则（在不违反质量边界的前提下优先执行）\n" + review.custom_rules.strip()
    issues, pending, governance = [], [], []
    # 分片大小与并发度是一对取舍：分片小则单批 JSON 输出短、不易触发网关
    # 120 秒上游超时，但分片小会成倍增加批次数。历史实现一度改成「10 行 +
    # 串行」，480 条用例产生 48 批串行调用，总耗时逼近 Celery 软时限。
    # 这里取 20 行 + 2 路并发，在两者之间取得平衡。
    chunk_size = TESTCASE_REVIEW_CHUNK_SIZE
    chunks = [rows[i:i + chunk_size] for i in range(0, len(rows), chunk_size)]
    review_timeout = max(30, min(int(config.request_timeout or 120), 600))
    attempts = _chunk_attempt_limit(config)
    workers = max(1, min(TESTCASE_REVIEW_MAX_WORKERS, len(chunks)))
    deadline = monotonic() + TESTCASE_REVIEW_TOTAL_BUDGET_SECONDS

    def review_one(index, chunk):
        # 底层 SDK 重试保持关闭：重试与退避统一由本层管理，单批最坏耗时才能
        # 按 attempts × review_timeout 估算，进而受总时间预算约束。
        llm = create_llm_instance(
            config,
            temperature=0.1,
            timeout=review_timeout,
            max_retries=0,
        )
        last_error = None
        for attempt in range(1, attempts + 1):
            try:
                return _review_chunk(llm, skill_prompt, chunk, review.business_context)
            except Exception as exc:
                last_error = exc
                logger.warning(
                    "用例审查分片失败 review=%s chunk=%s/%s attempt=%s/%s error=%s: %s",
                    review.id, index, len(chunks), attempt, attempts,
                    type(exc).__name__, exc,
                )
                if attempt < attempts:
                    sleep(_retry_delay(attempt))
        raise RuntimeError(
            f"第 {index}/{len(chunks)} 批模型调用失败（已尝试 {attempts} 次）："
            f"{type(last_error).__name__}: {last_error}"
        ) from last_error

    review.current_step = (
        f"Skill 审查 0/{len(chunks)}（并发 {workers}，"
        f"单批上限 {review_timeout} 秒）"
    )
    review.progress = 10
    review.save(update_fields=["current_step", "progress", "updated_at"])
    signature = hashlib.sha256(json.dumps({
        "rows": rows,
        "skill": skill_prompt,
        "business_context": review.business_context,
        "custom_rules": review.custom_rules,
        "chunk_size": chunk_size,
    }, ensure_ascii=False, sort_keys=True).encode("utf-8")).hexdigest()
    checkpoint = (review.summary or {}).get("_checkpoint") or {}
    ordered_results = {
        int(index): result
        for index, result in (checkpoint.get("results") or {}).items()
    } if checkpoint.get("signature") == signature else {}
    # 批次失败只记录、不中断整项任务：一次上游抖动不应让已跑完的批次白跑。
    uncovered = {}
    state_lock = threading.Lock()

    def persist_progress():
        processed = len(ordered_results) + len(uncovered)
        review.summary = {"_checkpoint": {
            "signature": signature,
            "total_chunks": len(chunks),
            "results": {str(key): value for key, value in ordered_results.items()},
            "uncovered": {str(key): value for key, value in uncovered.items()},
        }}
        review.current_step = (
            f"Skill 审查 {processed}/{len(chunks)}（并发 {workers}，"
            f"单批上限 {review_timeout} 秒）"
        )
        review.progress = 10 + int(processed / len(chunks) * 75)
        review.save(update_fields=["summary", "current_step", "progress", "updated_at"])

    waiting = [index for index in range(1, len(chunks) + 1) if index not in ordered_results]
    cursor = 0
    consecutive_failures = 0
    circuit_broken = False
    with ThreadPoolExecutor(max_workers=workers) as executor:
        while cursor < len(waiting):
            if monotonic() >= deadline:
                break
            wave = waiting[cursor:cursor + workers]
            cursor += len(wave)
            futures = {
                executor.submit(review_one, index, chunks[index - 1]): index
                for index in wave
            }
            for future in as_completed(futures):
                index = futures[future]
                try:
                    result = future.result()
                except Exception as exc:
                    with state_lock:
                        uncovered[index] = f"{type(exc).__name__}: {exc}"
                        consecutive_failures += 1
                        # 连续多批失败且迄今无任何成功批次，说明是网关整体不可用
                        # （而不是个别分片偶发失败）：提前终止，不必逐批耗尽重试。
                        if not ordered_results and consecutive_failures >= TESTCASE_REVIEW_CIRCUIT_BREAKER:
                            circuit_broken = True
                    logger.error(
                        "用例审查分片最终失败 review=%s chunk=%s/%s：%s",
                        review.id, index, len(chunks), exc,
                    )
                else:
                    with state_lock:
                        ordered_results[index] = result
                        consecutive_failures = 0
                with state_lock:
                    persist_progress()
            if circuit_broken:
                logger.error(
                    "用例审查连续 %s 批失败且无任何成功批次，判定模型网关不可用，提前终止 "
                    "review=%s（已送审 %s/%s 批）",
                    consecutive_failures, review.id, cursor, len(waiting),
                )
                break
    for index in waiting[cursor:]:
        uncovered.setdefault(
            index,
            "模型网关连续失败，已提前终止" if circuit_broken
            else f"总时间预算 {TESTCASE_REVIEW_TOTAL_BUDGET_SECONDS // 60} 分钟耗尽，该批未送审",
        )
    with state_lock:
        persist_progress()

    if not ordered_results:
        first_error = next(iter(uncovered.values()), "未知原因")
        raise RuntimeError(
            f"共 {len(uncovered)}/{len(chunks)} 批模型调用失败，未生成报告。首个错误：{first_error}"
        )

    # 按原始分片顺序汇总，保证报告顺序稳定可追溯。
    for index in range(1, len(chunks) + 1):
        result = ordered_results.get(index)
        if not result:
            continue
        issues.extend(result.get("issues") or [])
        pending.extend(result.get("pending_confirmations") or [])
        governance.extend(result.get("governance_suggestions") or [])
    review.current_step = "生成 Excel 报告"
    review.progress = 90
    review.summary = _write_report(review, rows, issues, pending, governance, uncovered, chunks)
    review.status = "completed"
    review.progress = 100
    review.current_step = (
        f"审查完成（{len(ordered_results)}/{len(chunks)} 批已覆盖）"
        if uncovered else "审查完成"
    )
    review.completed_at = timezone.now()
    review.error_message = ""
    review.save()
    return review.summary
