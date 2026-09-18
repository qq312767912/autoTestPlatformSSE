import os
import tempfile
from unittest.mock import patch

from django.contrib.auth import get_user_model
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase, override_settings
from openpyxl import Workbook, load_workbook
from rest_framework.test import APIClient

from projects.models import Project, ProjectMember
from testcases.models import TestCaseReview, TestCaseReviewLLMConfig
from testcases.review_service import _read_rows, _write_report
from skills.models import Skill


class TestCaseReviewApiTests(TestCase):
    def setUp(self):
        self.user = get_user_model().objects.create_user(username="reviewer", password="pass")
        self.project = Project.objects.create(name="Review Project", creator=self.user)
        ProjectMember.objects.create(project=self.project, user=self.user, role="member")
        # 审查任务的前置条件是「已配置专用 LLM」，否则创建会被 409 拦下。
        config = TestCaseReviewLLMConfig.objects.create(
            config_name="用例审查 LLM", name="review-model", api_url="http://model-service/v1",
        )
        config.set_api_key("secret")
        config.save()
        self.client = APIClient()
        self.client.force_authenticate(self.user)

    @override_settings(MEDIA_ROOT=tempfile.gettempdir())
    @patch("testcases.views.execute_testcase_review.delay")
    def test_upload_creates_async_review(self, delay):
        delay.return_value.id = "task-1"
        file = SimpleUploadedFile("cases.xlsx", b"fake", content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        response = self.client.post(
            f"/api/projects/{self.project.id}/testcase-reviews/",
            {"source_file": file, "business_context": "仅管理员可审批"},
            format="multipart",
        )
        self.assertEqual(response.status_code, 201)
        review = TestCaseReview.objects.get()
        self.assertEqual(review.business_context, "仅管理员可审批")
        self.assertEqual(review.celery_task_id, "task-1")
        self.assertTrue(response.data["source_url"].startswith("/media/"))
        delay.assert_called_once_with(review.id)
        review.delete()

    @patch("testcases.views.execute_testcase_review.delay")
    def test_upload_rejects_unsupported_file(self, delay):
        response = self.client.post(
            f"/api/projects/{self.project.id}/testcase-reviews/",
            {"source_file": SimpleUploadedFile("cases.txt", b"case")},
            format="multipart",
        )
        self.assertEqual(response.status_code, 400)
        delay.assert_not_called()

    @override_settings(MEDIA_ROOT=tempfile.gettempdir())
    @patch("testcases.views.execute_testcase_review.delay")
    def test_specified_skill_and_custom_rules_are_snapshotted(self, delay):
        delay.return_value.id = "task-2"
        skill = Skill.objects.create(
            project=self.project, creator=self.user, name="payment-review",
            description="支付专项审查", skill_content="# 支付审查\n检查金额和幂等性", is_active=True,
        )
        response = self.client.post(
            f"/api/projects/{self.project.id}/testcase-reviews/",
            {
                "source_file": SimpleUploadedFile("payment.xlsx", b"fake"),
                "review_mode": "specified", "selected_skill": skill.id,
                "custom_rules": "退款金额必须与原订单一致",
            },
            format="multipart",
        )
        self.assertEqual(response.status_code, 201)
        review = TestCaseReview.objects.get()
        self.assertEqual(review.skill_name, "payment-review")
        self.assertIn("检查金额和幂等性", review.skill_snapshot)
        self.assertEqual(review.custom_rules, "退款金额必须与原订单一致")
        review.delete()

    @patch("testcases.views.execute_testcase_review.delay")
    def test_specified_review_requires_skill(self, delay):
        response = self.client.post(
            f"/api/projects/{self.project.id}/testcase-reviews/",
            {"source_file": SimpleUploadedFile("cases.xlsx", b"fake"), "review_mode": "specified"},
            format="multipart",
        )
        self.assertEqual(response.status_code, 400)
        delay.assert_not_called()


class TestCaseReviewLLMGateTests(TestCase):
    """审查的前置条件：必须先由管理员配置并启用专用 LLM。"""

    def setUp(self):
        self.user = get_user_model().objects.create_user(username="gate-user", password="pass")
        self.project = Project.objects.create(name="Gate Project", creator=self.user)
        ProjectMember.objects.create(project=self.project, user=self.user, role="member")
        self.client = APIClient()
        self.client.force_authenticate(self.user)

    def _configure(self, api_key="secret"):
        config = TestCaseReviewLLMConfig.objects.create(
            config_name="用例审查 LLM", name="review-model", api_url="http://model-service/v1",
        )
        config.set_api_key(api_key)
        config.save()
        return config

    def _upload(self):
        with override_settings(MEDIA_ROOT=tempfile.gettempdir()):
            return self.client.post(
                f"/api/projects/{self.project.id}/testcase-reviews/",
                {"source_file": SimpleUploadedFile("cases.xlsx", b"fake")},
                format="multipart",
            )

    @patch("testcases.views.execute_testcase_review.delay")
    def test_upload_is_blocked_and_not_persisted_without_config(self, delay):
        # 前提自检：此刻确实没有任何可用配置，否则这条断言恒过。
        self.assertFalse(TestCaseReviewLLMConfig.objects.exists())
        response = self._upload()
        self.assertEqual(response.status_code, 409)
        # 关键：不能留下一条永远跑不起来的审查记录。
        self.assertFalse(TestCaseReview.objects.exists())
        delay.assert_not_called()

    @patch("testcases.views.execute_testcase_review.delay")
    def test_upload_allowed_once_config_exists(self, delay):
        delay.return_value.id = "task-gate"
        self._configure()
        response = self._upload()
        self.assertEqual(response.status_code, 201)
        delay.assert_called_once()
        TestCaseReview.objects.all().delete()

    @patch("testcases.views.execute_testcase_review.delay")
    def test_config_without_api_key_does_not_count_as_configured(self, delay):
        # 有配置但没密钥时依然按「未配置」处理，否则审查会在运行时才失败。
        self._configure(api_key="")
        response = self._upload()
        self.assertEqual(response.status_code, 409)
        delay.assert_not_called()


class TestCaseReviewLLMConfigPermissionTests(TestCase):
    """专用 LLM 配置只对平台管理员开放，普通成员连读都不允许。"""

    URL = "/api/testcases/review-llm-config/"

    def setUp(self):
        self.member = get_user_model().objects.create_user(username="plain-member", password="pass")
        self.admin = get_user_model().objects.create_user(
            username="platform-admin", password="pass", is_staff=True,
        )
        self.client = APIClient()

    def _list_as(self, user):
        self.client.force_authenticate(user)
        return self.client.get(self.URL)

    def test_regular_member_is_denied(self):
        self.assertEqual(self._list_as(self.member).status_code, 403)

    def test_admin_can_list(self):
        # 前提自检：管理员必须能读到，否则「拒绝普通成员」可能只是接口整体不通。
        self.assertEqual(self._list_as(self.admin).status_code, 200)

    def test_anonymous_is_rejected(self):
        self.client.force_authenticate(None)
        self.assertIn(self.client.get(self.URL).status_code, (401, 403))

    def test_admin_can_create_then_update_without_resending_key(self):
        self.client.force_authenticate(self.admin)
        created = self.client.post(self.URL, {
            "config_name": "用例审查 LLM", "name": "review-model",
            "api_url": "http://model-service/v1", "api_key": "secret",
            "request_timeout": 600, "max_retries": 2, "is_active": True,
        }, format="json")
        self.assertEqual(created.status_code, 201)
        self.assertTrue(created.data["has_api_key"])
        self.assertNotIn("api_key", created.data)

        config_id = created.data["id"]
        updated = self.client.patch(f"{self.URL}{config_id}/", {"name": "review-model-v2"}, format="json")
        self.assertEqual(updated.status_code, 200)
        # 留空表示保持原密钥，不得把已配置的密钥清掉。
        self.assertTrue(updated.data["has_api_key"])
        self.assertEqual(TestCaseReviewLLMConfig.objects.get(pk=config_id).name, "review-model-v2")

    def test_first_configuration_requires_api_key(self):
        self.client.force_authenticate(self.admin)
        response = self.client.post(self.URL, {
            "config_name": "用例审查 LLM", "name": "review-model",
            "api_url": "http://model-service/v1", "request_timeout": 600, "max_retries": 2,
        }, format="json")
        self.assertEqual(response.status_code, 400)
        self.assertIn("api_key", response.data)

    def test_out_of_range_timeout_is_rejected(self):
        self.client.force_authenticate(self.admin)
        response = self.client.post(self.URL, {
            "config_name": "用例审查 LLM", "name": "review-model",
            "api_url": "http://model-service/v1", "api_key": "secret", "request_timeout": 5,
        }, format="json")
        self.assertEqual(response.status_code, 400)
        self.assertIn("request_timeout", response.data)

    def test_configuration_stays_singleton(self):
        self.client.force_authenticate(self.admin)
        for name in ("first", "second"):
            response = self.client.post(self.URL, {
                "config_name": f"用例审查 LLM {name}", "name": "review-model",
                "api_url": "http://model-service/v1", "api_key": "secret",
            }, format="json")
            self.assertEqual(response.status_code, 201)
        self.assertEqual(TestCaseReviewLLMConfig.objects.count(), 1)


class TestCaseReviewReportTests(TestCase):
    def test_reads_xlsx_and_generates_traceable_report(self):
        with tempfile.TemporaryDirectory() as root, override_settings(MEDIA_ROOT=root):
            source = os.path.join(root, "cases.xlsx")
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "登录"
            sheet.append(["编号", "模块", "步骤", "预期"])
            sheet.append(["TC-1", "登录", "输入账号密码", "登录成功"])
            workbook.save(source)
            rows = _read_rows(source)
            self.assertEqual(rows[0]["sheet"], "登录")
            self.assertEqual(rows[0]["row"], 2)

            user = get_user_model().objects.create_user(username="report-user")
            project = Project.objects.create(name="Report Project", creator=user)
            review = TestCaseReview.objects.create(project=project, creator=user, source_name="cases.xlsx", source_file="source/cases.xlsx")
            summary = _write_report(review, rows, [{
                "sheet": "登录", "row": 2, "case_id": "TC-1", "module": "登录",
                "original": "登录成功", "severity": "中", "issue_type": "判定标准模糊",
                "description": "缺少可观察结果", "suggestion": "补充跳转页面和登录态",
                "judgement": "确定缺陷", "rewrite": "跳转首页且展示当前用户名",
            }], [], ["统一登录用例模板"])
            review.save()
            self.assertEqual(summary["medium"], 1)
            generated = load_workbook(review.report_file.path)
            self.assertIn("问题明细", generated.sheetnames)
            self.assertEqual(generated["问题明细"]["B2"].value, 2)
            self.assertEqual(generated["审查摘要"]["A1"].value, "测试用例质量审查报告")
            self.assertEqual(generated["审查摘要"]["A1"].fill.fgColor.rgb, "001D3F66")
            self.assertEqual(generated["问题明细"]["F2"].fill.fgColor.rgb, "00FFF2D6")
            self.assertFalse(generated["问题明细"].sheet_view.showGridLines)
            self.assertEqual(generated["问题明细"].freeze_panes, "A2")
            generated.close()
            review.delete()
