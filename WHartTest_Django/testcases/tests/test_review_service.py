import os
import tempfile
from types import SimpleNamespace
from unittest.mock import patch

from django.test import SimpleTestCase
from openpyxl import Workbook

from testcases import review_service
from testcases.review_service import (
    _chunk_attempt_limit,
    _read_rows,
    _retry_delay,
)


class TestCaseReviewReaderTests(SimpleTestCase):
    def test_recognizes_case_headers_and_filters_metadata_and_group_rows(self):
        with tempfile.TemporaryDirectory() as root:
            source = os.path.join(root, "cases.xlsx")
            workbook = Workbook()
            cover = workbook.active
            cover.title = "封面"
            cover.append(["系统测试用例"])

            cases = workbook.create_sheet("登录")
            cases.append(["编号", "模块", "步骤", "预期"])
            cases.append(["TC-1", "登录", "输入账号密码", "登录成功"])

            compatibility = workbook.create_sheet("兼容性")
            compatibility.append(["项目名称", "兼容性测试用例"])
            compatibility.append(["操作系统", "版本", "操作步骤", "期望输出"])
            compatibility.append(["Chrome 系列"])
            compatibility.append(["Windows 10", "64位", "打开页面", "页面布局完整"])
            workbook.save(source)

            rows = _read_rows(source)

        self.assertEqual(len(rows), 2)
        self.assertEqual(rows[0]["sheet"], "登录")
        self.assertEqual(rows[0]["row"], 2)
        self.assertEqual(rows[0]["columns"]["预期"], "登录成功")
        self.assertEqual(rows[1]["sheet"], "兼容性")
        self.assertEqual(rows[1]["row"], 4)


class ReviewRetryPolicyTests(SimpleTestCase):
    """重试策略：尊重 LLMConfig.max_retries，但有下限与上限，退避必须递增。"""

    def test_attempt_limit_honors_config_within_bounds(self):
        floor = review_service.TESTCASE_REVIEW_CHUNK_ATTEMPTS
        ceiling = review_service.TESTCASE_REVIEW_MAX_ATTEMPTS
        self.assertEqual(_chunk_attempt_limit(SimpleNamespace(max_retries=3)), 4)
        # 配置为 0（界面语义是「禁用重试」）不得让分片退化为零容错。
        self.assertEqual(_chunk_attempt_limit(SimpleNamespace(max_retries=0)), floor)
        self.assertEqual(_chunk_attempt_limit(SimpleNamespace(max_retries=99)), ceiling)
        # 配置缺失或非法时回落到下限，而不是抛错。
        self.assertEqual(_chunk_attempt_limit(SimpleNamespace(max_retries=None)), floor)
        self.assertEqual(_chunk_attempt_limit(SimpleNamespace(max_retries="bad")), floor)

    def test_retry_delay_is_exponential_and_capped(self):
        self.assertEqual([_retry_delay(attempt) for attempt in range(1, 7)],
                         [2, 4, 8, 16, 30, 30])


class TestCaseReviewLLMConfigResolutionTests(SimpleTestCase):
    """审查只认专用配置：缺失或没密钥必须失败，不得回退平台通用 LLM。"""

    def _resolve(self, config):
        with patch.object(review_service, "TestCaseReviewLLMConfig", SimpleNamespace(
                objects=SimpleNamespace(filter=lambda **kwargs: SimpleNamespace(first=lambda: config)))):
            return review_service._get_testcase_review_llm_config()

    def test_module_has_no_platform_llm_entry_point(self):
        # 前提自检：模块里若还留着平台通用 LLMConfig，实现就可能悄悄回退，
        # 下面两条断言也就失去了意义。
        self.assertFalse(hasattr(review_service, "LLMConfig"))

    def test_missing_config_raises_instead_of_falling_back(self):
        with self.assertRaises(ValueError) as raised:
            self._resolve(None)
        self.assertIn("用例审查专用 LLM", str(raised.exception))

    def test_config_without_api_key_raises(self):
        with self.assertRaises(ValueError) as raised:
            self._resolve(SimpleNamespace(api_key=""))
        self.assertIn("API Key", str(raised.exception))

    def test_active_config_is_returned(self):
        config = SimpleNamespace(api_key="k", request_timeout=100, max_retries=1)
        self.assertIs(self._resolve(config), config)


class ReviewChunkFallbackTests(SimpleTestCase):
    """分片失败只降级、不中断整项；网关整体故障时提前熔断。"""

    TOTAL_ROWS = 40

    def _run(self, *, failing_rows=(), max_retries=3, budget=None, total_rows=None):
        rows = [
            {"sheet": "登录", "row": index + 2, "cells": ["TC", "步骤", "预期"],
             "columns": {"编号": "TC", "步骤": "步骤", "预期": "预期"}}
            for index in range(total_rows or self.TOTAL_ROWS)
        ]
        calls, sleeps = [], []
        review = SimpleNamespace(
            id=1, status="pending", progress=0, current_step="", business_context="",
            custom_rules="", skill_snapshot="SKILL", selected_skill=None, summary={},
            error_message="", source_name="cases.xlsx", skill_name="test-case-clarity-review",
            source_file=SimpleNamespace(path="/tmp/cases.xlsx"),
            report_file=SimpleNamespace(save=lambda *args, **kwargs: None),
            save=lambda **kwargs: None,
            get_review_mode_display=lambda: "通用审查",
        )
        config = SimpleNamespace(request_timeout=120, max_retries=max_retries, api_key="test-key")

        def fake_review_chunk(llm, skill_prompt, chunk_rows, business_context):
            calls.append(chunk_rows[0]["row"])
            if chunk_rows[0]["row"] in failing_rows:
                raise RuntimeError("Error code: 502 - upstream_unavailable")
            return {"issues": [], "pending_confirmations": [], "governance_suggestions": []}

        patches = [
            patch.object(review_service, "_read_rows", return_value=rows),
            patch.object(review_service, "TestCaseReview", SimpleNamespace(
                objects=SimpleNamespace(get=lambda **kwargs: review))),
            patch.object(review_service, "TestCaseReviewLLMConfig", SimpleNamespace(
                objects=SimpleNamespace(filter=lambda **kwargs: SimpleNamespace(first=lambda: config)))),
            patch.object(review_service, "create_llm_instance", return_value=object()),
            patch.object(review_service, "_review_chunk", side_effect=fake_review_chunk),
            patch.object(review_service, "sleep", side_effect=lambda seconds: sleeps.append(seconds)),
        ]
        if budget is not None:
            patches.append(patch.object(review_service, "TESTCASE_REVIEW_TOTAL_BUDGET_SECONDS", budget))
        for item in patches:
            item.start()
            self.addCleanup(item.stop)

        error, summary = None, None
        try:
            summary = review_service.run_testcase_review(1)
        except Exception as exc:
            error = exc
        chunk_size = review_service.TESTCASE_REVIEW_CHUNK_SIZE
        return SimpleNamespace(
            error=error, summary=summary, calls=calls, sleeps=sleeps, review=review,
            total_chunks=-(-len(rows) // chunk_size), chunk_size=chunk_size,
        )

    def test_single_failing_chunk_is_downgraded_not_fatal(self):
        result = self._run(failing_rows={2})
        self.assertIsNone(result.error)
        self.assertIsNotNone(result.summary)
        # 失败批次被标注为未覆盖，其余批次结果照常产出。
        self.assertEqual(result.summary["uncovered_chunks"], 1)
        self.assertEqual(result.summary["uncovered_rows"], result.chunk_size)
        self.assertEqual(result.summary["total_chunks"], result.total_chunks)

    def test_failing_chunk_retries_with_exponential_backoff(self):
        result = self._run(failing_rows={2}, max_retries=3)
        attempts = sum(1 for row in result.calls if row == 2)
        self.assertEqual(attempts, 4)                      # max_retries + 1
        self.assertEqual(result.sleeps[:3], [2, 4, 8])     # 递增退避

    def test_all_chunks_failing_raises_without_report(self):
        result = self._run(failing_rows=set(range(2, 2 + self.TOTAL_ROWS)))
        self.assertIsNotNone(result.error)
        self.assertIn("未生成报告", str(result.error))
        self.assertIsNone(result.summary)

    def test_consecutive_failures_break_early(self):
        # 需要足够多的批次才能观察到熔断：熔断阈值是「连续失败批数」，
        # 批数少于阈值时必然全部送审，测不出提前终止。
        total_rows = 10 * review_service.TESTCASE_REVIEW_CHUNK_SIZE
        result = self._run(failing_rows=set(range(2, 2 + total_rows)), total_rows=total_rows)
        sent = len(set(result.calls))
        self.assertGreater(result.total_chunks, 2 * review_service.TESTCASE_REVIEW_CIRCUIT_BREAKER)
        # 网关整体不可用时不必逐批耗尽重试：送审批次应远小于总批数。
        self.assertLessEqual(sent, 2 * review_service.TESTCASE_REVIEW_CIRCUIT_BREAKER)
        self.assertLess(sent, result.total_chunks)

    def test_budget_exhaustion_marks_all_chunks_uncovered(self):
        result = self._run(budget=0)
        self.assertEqual(result.calls, [])
        self.assertIsNotNone(result.error)
        uncovered = result.review.summary["_checkpoint"]["uncovered"]
        self.assertEqual(len(uncovered), result.total_chunks)
