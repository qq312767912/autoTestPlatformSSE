"""T07：方案阶段适配器与平台派生产物。

验收盯四件事，每一件都对应一种"页面会给出看起来正常的错数据"的失败：

- **确定性**：同一输入重复生成，图谱/摘要逐字节一致、报告的行与统计一致。
  不一致会让"这条摘要变了"无法区分"产出变了"还是"只是又跑了一次生成器"。
- **引用状态**：只有确定性定位验证通过才标 ``verified``；定位不到必须是 ``invalid``，
  不能因为"引用了"就当成可信证据。
- **业务产物隔离**：平台派生物不覆盖 Skill 的主产物（文件名撞车直接拒绝）。
- **登记留痕**：生成器版本与源产出哈希必须落库，否则"结论为什么变了"永远查不出来。
"""
from __future__ import annotations

import hashlib
import json
import tempfile
from pathlib import Path
from unittest import mock

from django.test import SimpleTestCase, override_settings

from knowledge_evolution import derived_artifacts as da
from knowledge_evolution.evidence_graph import GRAPH_FILENAME
from knowledge_evolution.quality_summary import QUALITY_SUMMARY_FILENAME
from knowledge_evolution.review_reports import (
    HUMAN_COLUMNS,
    REVIEW_REPORT_FILENAME,
    SHEET_ITEM_ROWS,
    report_statistics_from_workbook,
)
from knowledge_evolution.stage_outputs import (
    StageOutputAdapter,
    StageProtocolError,
    TestPlanOutputAdapter,
    get_adapter,
    validate_stage_result,
)
from knowledge_evolution.tests_t09_t13 import SkillHubBaseTests
from knowledge_evolution.workflow_models import StageDerivedArtifact

#: 派生物会真实落盘，绝不能写进项目数据目录。
TEST_MEDIA_ROOT = tempfile.mkdtemp(prefix="wharttest-t07-")

STAGE = "test_plan_generation"
QUOTE = "非白名单证券账户不得进入后续投票处理流程。"


def _quote_hash() -> str:
    return hashlib.sha256(QUOTE.encode("utf-8")).hexdigest()


def _payload(*, start_offset: int = 40, content_hash: str | None = None) -> dict:
    return {
        "schema_version": "stage-result/v1",
        "stage": STAGE,
        "status": "completed",
        "primary_artifacts": [{"name": "test_plan.xlsx", "kind": "xlsx"}],
        "requirements": [{"id": "REQ-12", "title": "非白名单账户校验"}],
        "items": [
            {
                "id": "TP-018",
                "title": "非白名单账户反向校验",
                "requirement_ids": ["REQ-12"],
                "evidence_ids": ["EV-001"],
                "scenario_type": "negative",
                "priority": "P0",
                "expected": "非白名单账户被拒绝并记录审计日志",
            }
        ],
        "evidence": [
            {
                "id": "EV-001",
                "document_id": "doc-21",
                "document_version": "v3",
                "chunk_id": "chunk-108",
                "quote": QUOTE,
                "start_offset": start_offset,
                "content_hash": content_hash if content_hash is not None else f"sha256:{_quote_hash()}",
            }
        ],
        "decision_summary": [{"decision": "按反向场景覆盖", "reason": "业务规则明确禁止"}],
        "uncertainties": ["白名单更新时效待确认"],
    }


def _corpus(*, text: str | None = None, start_offset: int = 40) -> dict:
    """原文索引：``quote`` 位于片段第 0 位，绝对偏移 = 片段起始偏移。"""
    return {
        "documents": [{
            "document_id": "doc-21",
            "document_version": "v3",
            "chunks": [{
                "chunk_id": "chunk-108",
                "start_offset": start_offset,
                "text": text if text is not None else QUOTE + "（后略）",
            }],
        }]
    }


def _stage_result(payload: dict | None = None):
    validation = validate_stage_result(payload or _payload(), expected_stage=STAGE)
    assert validation.ok, validation.as_dict()
    return TestPlanOutputAdapter().parse(validation.payload or {})


# ---------------------------------------------------------------------------
# 适配器接口
# ---------------------------------------------------------------------------


class AdapterInterfaceTests(SimpleTestCase):
    def test_base_adapter_methods_are_not_implemented(self):
        adapter = StageOutputAdapter()
        for name in ("parse", "build_review_report", "build_evidence_graph", "build_quality_summary"):
            with self.assertRaises(NotImplementedError):
                getattr(adapter, name)(None)

    def test_registry_returns_none_for_unimplemented_stage(self):
        # 返回空壳会让人以为"这个阶段已支持"，显式 None 才诚实。
        self.assertIsNone(get_adapter("testcase_generation"))
        self.assertIsNone(get_adapter(""))
        adapter = get_adapter(STAGE)
        self.assertIsInstance(adapter, TestPlanOutputAdapter)
        self.assertEqual(adapter.stage, STAGE)

    def test_parse_rejects_invalid_payload_with_protocol_error(self):
        adapter = TestPlanOutputAdapter()
        with self.assertRaises(StageProtocolError) as ctx:
            adapter.parse({"stage": STAGE})
        self.assertFalse(ctx.exception.validation.ok)
        self.assertTrue(ctx.exception.validation.issues)


# ---------------------------------------------------------------------------
# 内容生成（纯函数层）
# ---------------------------------------------------------------------------


class DerivedContentTests(SimpleTestCase):
    def test_builds_three_registered_kinds(self):
        contents = da.build_derived_contents(_stage_result(), corpus=_corpus())
        self.assertEqual(
            set(contents), {da.KIND_EVIDENCE_GRAPH, da.KIND_QUALITY_SUMMARY, da.KIND_REVIEW_REPORT},
        )
        self.assertEqual(set(da.KIND_FILENAMES), set(contents))
        # 三种内容的文件名必须互不相同，也不与业务主产物同名。
        names = set(da.KIND_FILENAMES.values())
        self.assertEqual(len(names), 3)
        self.assertFalse(names & da.BUSINESS_ARTIFACT_FILENAMES)

    def test_json_contents_are_byte_identical_across_runs(self):
        stage_result = _stage_result()
        first = da.build_derived_contents(stage_result, corpus=_corpus())
        second = da.build_derived_contents(stage_result, corpus=_corpus())
        self.assertEqual(first[da.KIND_EVIDENCE_GRAPH], second[da.KIND_EVIDENCE_GRAPH])
        self.assertEqual(first[da.KIND_QUALITY_SUMMARY], second[da.KIND_QUALITY_SUMMARY])

    def test_report_rows_and_statistics_are_stable_across_runs(self):
        """xlsx 字节含 zip 时间戳，验收的是行与统计一致（见 review_reports 注明）。"""
        stage_result = _stage_result()
        first = report_statistics_from_workbook(
            da.build_derived_contents(stage_result, corpus=_corpus())[da.KIND_REVIEW_REPORT]
        )
        second = report_statistics_from_workbook(
            da.build_derived_contents(stage_result, corpus=_corpus())[da.KIND_REVIEW_REPORT]
        )
        self.assertEqual(first, second)

    def test_human_columns_start_blank(self):
        """人工列初始必须为空——默认值会被当成已填，把"没人看过"伪装成"审核通过"。"""
        from io import BytesIO

        from openpyxl import load_workbook

        payload = da.build_derived_contents(_stage_result(), corpus=_corpus())[da.KIND_REVIEW_REPORT]
        sheet = load_workbook(BytesIO(payload))[SHEET_ITEM_ROWS]
        rows = list(sheet.iter_rows(values_only=True))
        headers = list(rows[0])
        for column in HUMAN_COLUMNS:
            index = headers.index(column)
            for row in rows[1:]:
                self.assertEqual(str(row[index] or ""), "")

    def test_verified_only_after_deterministic_check(self):
        contents = da.build_derived_contents(_stage_result(), corpus=_corpus())
        graph = json.loads(contents[da.KIND_EVIDENCE_GRAPH].decode("utf-8"))
        self.assertEqual(graph["evidence_status"]["EV-001"]["status"], "verified")
        self.assertEqual(graph["summary"]["evidence_status"], {"verified": 1})

    def test_unlocatable_quote_is_invalid_not_silently_trusted(self):
        contents = da.build_derived_contents(
            _stage_result(), corpus=_corpus(text="完全不相干的另一段原文"),
        )
        graph = json.loads(contents[da.KIND_EVIDENCE_GRAPH].decode("utf-8"))
        self.assertEqual(graph["evidence_status"]["EV-001"]["status"], "invalid")
        quality = json.loads(contents[da.KIND_QUALITY_SUMMARY].decode("utf-8"))
        self.assertEqual(quality["evidence"]["by_status"], {"invalid": 1})
        self.assertNotIn("verified", quality["evidence"]["by_status"])

    def test_offset_mismatch_is_invalid(self):
        """引用的那句话没错、定位信息却是上一版的——必须判失败而不是将就通过。"""
        contents = da.build_derived_contents(
            _stage_result(_payload(start_offset=999)), corpus=_corpus(),
        )
        graph = json.loads(contents[da.KIND_EVIDENCE_GRAPH].decode("utf-8"))
        self.assertEqual(graph["evidence_status"]["EV-001"]["status"], "invalid")
        self.assertEqual(
            graph["evidence_status"]["EV-001"]["issues"][0]["code"], "offset_mismatch",
        )

    def test_declared_status_when_corpus_missing(self):
        """运行时拿不到索引时是 ``declared``：'验不了'与'验错了'必须分开。"""
        contents = da.build_derived_contents(_stage_result())
        graph = json.loads(contents[da.KIND_EVIDENCE_GRAPH].decode("utf-8"))
        self.assertEqual(graph["evidence_status"]["EV-001"]["status"], "declared")

    def test_source_hash_differs_between_bytes_and_object(self):
        payload = _payload()
        as_object = da.source_output_hash(payload)
        as_bytes = da.source_output_hash(json.dumps(payload).encode("utf-8"))
        self.assertNotEqual(as_object, as_bytes)
        self.assertEqual(len(as_object), 64)


# ---------------------------------------------------------------------------
# 落盘与登记（需要数据库）
# ---------------------------------------------------------------------------


@override_settings(MEDIA_ROOT=TEST_MEDIA_ROOT)
class DerivedRegistrationTests(SkillHubBaseTests):
    def setUp(self):
        super().setUp()
        self.output = self.make_output(task_type=STAGE, workflow_id="wf-t07")

    def _register(self, payload: dict | None = None, **kwargs):
        return da.register_for_output(
            self.output,
            stage_result_payload=payload if payload is not None else _payload(),
            actor=self.lead,
            **kwargs,
        )

    def test_registers_three_artifacts_with_version_and_source_hash(self):
        artifacts = self._register(corpus=_corpus())
        self.assertEqual(len(artifacts), 3)
        kinds = {item.kind for item in artifacts}
        self.assertEqual(kinds, set(da.KIND_FILENAMES))
        expected_source_hash = da.source_output_hash(_payload())
        for item in artifacts:
            self.assertEqual(item.project_id, self.output.project_id)
            self.assertEqual(item.stage, STAGE)
            self.assertEqual(item.workflow_id, "wf-t07")
            self.assertEqual(item.level, "L3")
            self.assertEqual(item.source_output_hash, expected_source_hash)
            self.assertEqual(
                item.generator_version, da.KIND_GENERATOR_VERSIONS[item.kind],
            )
            self.assertTrue(da.absolute_path(item).exists())
            self.assertEqual(da.read_artifact(item), da.absolute_path(item).read_bytes())
            self.assertEqual(item.content_hash, da.hash_bytes(da.read_artifact(item)))
        self.assertEqual(StageDerivedArtifact.objects.filter(output=self.output).count(), 3)

    def test_media_root_relative_path_is_recorded(self):
        artifact = next(
            item for item in self._register(corpus=_corpus())
            if item.kind == da.KIND_EVIDENCE_GRAPH
        )
        self.assertTrue(artifact.path.startswith(f"{da.DERIVED_ROOT_NAME}/"))
        self.assertFalse(Path(artifact.path).is_absolute())
        self.assertEqual(artifact.filename, GRAPH_FILENAME)

    def test_repeated_registration_is_idempotent(self):
        first = {item.kind: item.content_hash for item in self._register(corpus=_corpus())}
        second = {item.kind: item.content_hash for item in self._register(corpus=_corpus())}
        self.assertEqual(first, second)
        self.assertEqual(StageDerivedArtifact.objects.filter(output=self.output).count(), 3)
        for artifact in StageDerivedArtifact.objects.filter(output=self.output):
            self.assertEqual(len(artifact.metadata["history"]), 1)

    def test_source_change_appends_history_without_duplicating_rows(self):
        self._register(corpus=_corpus())
        artifacts = da.register_for_output(
            self.output, stage_result_payload=_payload(), actor=self.lead,
            corpus=_corpus(), source_hash="b" * 64,
        )
        self.assertEqual(StageDerivedArtifact.objects.filter(output=self.output).count(), 3)
        for artifact in artifacts:
            self.assertEqual(artifact.source_output_hash, "b" * 64)
            self.assertEqual(len(artifact.metadata["history"]), 2)
            self.assertEqual(artifact.metadata["history"][0]["source_output_hash"] != "b" * 64, True)

    def test_history_is_capped(self):
        for index in range(da.HISTORY_LIMIT + 5):
            da.register_for_output(
                self.output, stage_result_payload=_payload(), actor=self.lead,
                corpus=_corpus(), source_hash=f"{index:064d}",
            )
        artifact = StageDerivedArtifact.objects.get(
            output=self.output, kind=da.KIND_EVIDENCE_GRAPH,
        )
        self.assertEqual(len(artifact.metadata["history"]), da.HISTORY_LIMIT)

    def test_derived_files_never_overwrite_business_artifacts(self):
        business_dir = Path(settings_media_root()) / "skill_runtime" / str(self.output.pk)
        business_dir.mkdir(parents=True, exist_ok=True)
        business_file = business_dir / "test_plan.xlsx"
        business_bytes = "原始业务方案".encode("utf-8")
        business_file.write_bytes(business_bytes)

        da.register_for_output(
            self.output, stage_result_payload=_payload(), actor=self.lead, corpus=_corpus(),
        )

        self.assertEqual(business_file.read_bytes(), business_bytes)
        derived_dir = da.output_directory(self.output)
        self.assertNotEqual(derived_dir, business_dir)
        self.assertFalse(
            set(path.name for path in derived_dir.iterdir()) & set(da.BUSINESS_ARTIFACT_FILENAMES)
        )

    def test_business_artifact_filename_is_refused(self):
        with self.assertRaises(ValueError):
            da.assert_not_business_artifact(Path("/tmp") / "stage_result.json")

    def test_unknown_kind_is_rejected(self):
        with self.assertRaises(ValueError):
            da.relative_path_for(self.output, "not_a_kind")

    def test_skips_unimplemented_stage(self):
        other = self.make_output(task_type="testcase_generation", workflow_id="wf-t07b")
        self.assertEqual(
            da.register_for_output(other, stage_result_payload=_payload(), actor=self.lead), [],
        )
        self.assertFalse(StageDerivedArtifact.objects.filter(output=other).exists())

    def test_skips_when_payload_absent(self):
        self.assertEqual(da.register_for_output(self.output, actor=self.lead), [])
        self.assertFalse(StageDerivedArtifact.objects.filter(output=self.output).exists())

    def test_protocol_failure_is_surfaced_not_swallowed(self):
        """结构化失败与业务生成失败必须是两个标记，所以这里必须往上抛。"""
        broken = _payload()
        broken.pop("primary_artifacts")
        with self.assertRaises(StageProtocolError):
            self._register(payload=broken)
        self.assertFalse(StageDerivedArtifact.objects.filter(output=self.output).exists())

    def test_storage_failure_is_best_effort_by_default(self):
        """观测面故障不能让业务产出一起丢：默认吞掉并返回空列表。"""
        with mock.patch.object(da, "output_directory", side_effect=OSError("磁盘满了")):
            self.assertEqual(self._register(), [])

    def test_storage_failure_can_be_strict(self):
        with mock.patch.object(da, "output_directory", side_effect=OSError("磁盘满了")):
            with self.assertRaises(OSError):
                self._register(best_effort=False)

    def test_list_for_output_returns_registered_rows(self):
        self._register(corpus=_corpus())
        kinds = [item.kind for item in da.list_for_output(self.output)]
        self.assertEqual(kinds, sorted(da.KIND_FILENAMES))

    def test_output_hash_changed_detects_replaced_source(self):
        artifact = next(
            item for item in self._register(corpus=_corpus())
            if item.kind == da.KIND_QUALITY_SUMMARY
        )
        self.assertFalse(da.output_hash_changed(artifact, self.output))
        self.output.metadata = {
            "protocol": {"stage_result_hash": "c" * 64},
        }
        self.output.save(update_fields=["metadata"])
        self.assertTrue(da.output_hash_changed(artifact, self.output))


def settings_media_root() -> str:
    from django.conf import settings

    return str(settings.MEDIA_ROOT)
