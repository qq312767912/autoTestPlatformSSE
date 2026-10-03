"""全链路自进化闭环：采纳率解析公共件（R3）+ 阶段报告下载（R1）+ 阶段反馈（R2）。

本文件钉住四件容易在真实使用里出错的事：

1. **采纳率的三种写法都要读对**（``0.85`` / ``85`` / ``"85%"``），
   且公式无缓存值不能被当成 0 —— 那会把"模板里的公式没算"表现成"这一版质量极差"。
2. **缺「采纳率」要报错、且报的是"文件不对"**，不能静默记成"未评分"：
   拿一份 Word 报告当依据，用户会以为记录成功了。
3. **采纳率不拦截**：低于参考线照常入库，只在返回值里标 ``below_reference``。
4. **反馈与派生解耦**：上传反馈不许动 Skill 包、不许落归因。

判据来源：``specs/evolution-assisted-loop/{requirements,design,tasks}.md``。
"""
from __future__ import annotations

import hashlib
import tempfile
from io import BytesIO
from pathlib import Path

from django.core.exceptions import ValidationError
from django.test import override_settings

from knowledge_evolution.models import FeedbackEvent, GenerationOutput
from knowledge_evolution.operations import WORKFLOW_STAGE_ORDER, WorkflowGateService
from knowledge_evolution.report_parsing import (
    DEFAULT_ACCEPTANCE_REFERENCE,
    locate_acceptance,
    read_acceptance_from_workbook,
)
from knowledge_evolution.tests_t09_t13 import TEST_MEDIA_ROOT, SkillHubBaseTests
from knowledge_evolution.trace_models import FailureAttribution
from knowledge_evolution.workflow_feedback import WorkflowStageFeedbackService
from skills.models import SkillVersion

PLAN = WORKFLOW_STAGE_ORDER[0]
BASE = "/api/knowledge-evolution/operations/"
ARTIFACT_URL = f"{BASE}workflow-stage-artifact/"
FEEDBACK_URL = f"{BASE}workflow-stage-feedback/"


def make_stage_report(rate=None, *, label: str = "采纳率", sheet_name: str = "汇总") -> bytes:
    """造一份阶段报告：最后一个 Sheet 里带「采纳率」标签单元格。

    用 openpyxl 真写文件而不是塞字典：这一层的风险恰恰在"真实 xlsx 里
    数字与文本混着长什么样"，用内存对象测会把数字型 85 与文本型 "85%" 的差异绕过去。
    """
    from openpyxl import Workbook

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = sheet_name
    sheet.append(["阶段", "结论"])
    sheet.append([PLAN, "已完成"])
    if rate is not None:
        sheet.append([label, rate])
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def make_formula_report() -> bytes:
    """采纳率是公式且**没有缓存值** —— 平台自己导出的报告正是这种形态。"""
    from openpyxl import Workbook

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "汇总"
    sheet.append(["采纳率", "=B1/C1"])
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def detail_text(response) -> str:
    """把 DRF 的错误体拍成一句话。

    ``raise ValidationError("...")``（字符串）出来的 ``detail`` 是 list，
    ``Response({"detail": ...})`` 出来的是 str —— 两种形状都会碰到，
    断言不该绑在某一种上。
    """
    detail = response.data["detail"] if isinstance(response.data, dict) else response.data
    if isinstance(detail, (list, tuple)):
        return "；".join(str(item) for item in detail)
    return str(detail)


class AcceptanceParsingTests(SkillHubBaseTests):
    """采纳率解析共同件：三种写法、公式、缺标签、越界。"""

    @staticmethod
    def _last_sheet(data: bytes):
        from openpyxl import load_workbook

        return load_workbook(BytesIO(data), data_only=True, read_only=True).worksheets[-1]

    def test_reads_three_number_spellings_as_the_same_ratio(self):
        """``0.85`` / ``85`` / ``"85%"`` 必须读出同一个比率。

        三种写法在真实报告里都会出现（人工填、公式算、文本粘贴），
        只认一种会让同一条结论在不同报告里得出不同分数。
        """
        for raw in (0.85, 85, "85%"):
            with self.subTest(raw=raw):
                sheet = self._last_sheet(make_stage_report(raw))
                found, value = locate_acceptance(sheet)
                self.assertTrue(found)
                self.assertAlmostEqual(value, 0.85, places=6)

    def test_formula_without_cached_value_is_not_zero(self):
        """公式无缓存值 → ``(True, None)``，**不是 0**。

        平台导出的报告采纳率正是公式；把 None 当 0 会让"公式没算"表现成"质量极差"，
        两种结论天差地别。
        """
        sheet = self._last_sheet(make_formula_report())
        found, value = locate_acceptance(sheet)
        self.assertTrue(found)
        self.assertIsNone(value)

    def test_out_of_range_rate_raises_instead_of_silently_clamping(self):
        """越界值报错。静默截断会让"模板写错了"表现成"这一版质量很差"。"""
        sheet = self._last_sheet(make_stage_report(150))
        with self.assertRaises(ValidationError):
            locate_acceptance(sheet)

    def test_missing_label_respects_the_required_flag(self):
        """缺标签：``required=True`` 报错、``required=False`` 返回 None。"""
        from openpyxl import load_workbook

        data = make_stage_report(None)  # 有表但没有「采纳率」标签
        workbook = load_workbook(BytesIO(data), data_only=True, read_only=True)

        value, _sheet = read_acceptance_from_workbook(workbook, required=False)
        self.assertIsNone(value)

        with self.assertRaises(ValidationError) as ctx:
            read_acceptance_from_workbook(workbook, required=True)
        self.assertIn("采纳率", str(ctx.exception))

    def test_reference_line_is_the_single_public_constant(self):
        """参考线只有一处真值：``70``。别处再写一遍就会分叉。"""
        self.assertEqual(DEFAULT_ACCEPTANCE_REFERENCE, 70.0)


@override_settings(MEDIA_ROOT=TEST_MEDIA_ROOT)
class StageArtifactTests(SkillHubBaseTests):
    """「下载报告」：登记产物优先，模板回落兜底。"""

    def setUp(self):
        super().setUp()
        from rest_framework.test import APIClient

        self.client = APIClient()
        self.client.force_authenticate(self.executor)

    def _publish(self, *, content="## 方案正文\n覆盖登录与支付", artifacts=None,
                 workflow_id="wf-art"):
        _skill, version = self.make_skill_version(name="plan-skill", stage=PLAN)
        output = self.make_output(
            task_type=PLAN, skill_version=version, workflow_id=workflow_id,
        )
        output.content = content
        if artifacts is not None:
            output.metadata = {**(output.metadata or {}), "artifacts": artifacts}
        output.save(update_fields=["content", "metadata"])
        WorkflowGateService.register_output(output)
        return output

    def _download(self, *, workflow_id="wf-art", stage=PLAN):
        return self.client.get(
            ARTIFACT_URL,
            {"project": self.project.id, "workflow_id": workflow_id, "stage": stage},
        )

    def test_registered_artifact_is_served_with_its_sha256(self):
        """有登记产物时下发的就是那个文件，且带 sha256 供前端核验。"""
        payload = "阶段报告二进制内容".encode("utf-8")
        path = Path(TEST_MEDIA_ROOT) / "skill_runtime" / "artifacts" / "plan-report.xlsx"
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_bytes(payload)
        digest = hashlib.sha256(payload).hexdigest()

        self._publish(artifacts=[{
            "key": "report", "name": "方案阶段报告_回归.xlsx",
            "path": str(path), "sha256": digest, "size": len(payload),
        }])

        response = self._download()
        self.assertEqual(response.status_code, 200)
        self.assertEqual(b"".join(response.streaming_content) if response.streaming else response.content, payload)
        self.assertEqual(response["X-Artifact-Source"], "registered")
        self.assertEqual(response["X-Artifact-Sha256"], digest)
        # 中文名必须走 RFC 5987，否则浏览器按 latin-1 解出乱码文件名。
        self.assertIn("filename*=UTF-8''", response["Content-Disposition"])

    def test_path_relative_to_media_root_is_resolved(self):
        """登记的是相对 ``MEDIA_ROOT`` 的路径（Django FileField 的存法）也要能找到。"""
        path = Path(TEST_MEDIA_ROOT) / "skill_runtime" / "artifacts" / "rel.txt"
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_text("相对路径产物", encoding="utf-8")

        self._publish(artifacts=[{
            "key": "report", "name": "rel.txt",
            "path": "skill_runtime/artifacts/rel.txt", "sha256": "",
        }])

        response = self._download()
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response["X-Artifact-Source"], "registered")

    def test_missing_registered_file_falls_back_to_template(self):
        """登记了但文件不在盘上 → 回落成模板，而不是 500。

        "文件不存在"等价于"没登记"：先给按钮再报 500，是把"没登记"说成"平台坏了"。
        """
        self._publish(
            content="## 方案正文\n回落也要能读到这句话",
            artifacts=[{
                "key": "report", "name": "gone.xlsx",
                "path": str(Path(TEST_MEDIA_ROOT) / "not-there.xlsx"), "sha256": "",
            }],
        )

        response = self._download()
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response["X-Artifact-Source"], "fallback")
        self.assertIn("text/markdown", response["Content-Type"])
        body = response.content.decode("utf-8")
        self.assertIn("回落也要能读到这句话", body)
        # 回落文件必须自我说明来路，避免被当成正式报告。
        self.assertIn("回落生成", body)

    def test_stage_without_output_is_400_not_empty_file(self):
        """未产出阶段明确说"无产出"，而不是回一个 0 字节附件。

        空文件会被读成"报告是空的"，而真相是"这一阶段还没跑"。
        """
        response = self._download(workflow_id="wf-never")
        self.assertEqual(response.status_code, 400)
        self.assertIn("暂无产出", detail_text(response))

    def test_unknown_stage_is_rejected(self):
        response = self._download(stage="no_such_stage")
        self.assertEqual(response.status_code, 400)

    def test_other_project_output_is_not_reachable(self):
        """三元定位：同 workflow_id 但别的项目的产出读不到。

        只按 workflow_id 取会让"越权读取"退化成"取决于 id 是否被猜中"。
        """
        payload = "别的项目的报告".encode("utf-8")
        path = Path(TEST_MEDIA_ROOT) / "other" / "report.txt"
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_bytes(payload)

        _skill, version = self.make_skill_version(name="other-skill", stage=PLAN,
                                                  project=self.other_project)
        foreign = self.make_output(task_type=PLAN, skill_version=version,
                                   project=self.other_project, workflow_id="wf-shared")
        foreign.metadata = {**(foreign.metadata or {}), "artifacts": [{
            "key": "report", "name": "report.txt", "path": str(path), "sha256": "",
        }]}
        foreign.save(update_fields=["metadata"])

        # 本项目的同名流程：查不到产出 → 400，绝不能拿到别项目的文件。
        response = self._download(workflow_id="wf-shared")
        self.assertEqual(response.status_code, 400)

    def test_workflow_status_exposes_artifact_availability(self):
        """阶段卡片靠 ``output.artifact`` 决定渲不渲染「下载报告」。"""
        self._publish(content="正文", artifacts=None)
        response = self.client.get(
            f"{BASE}workflow-status/",
            {"project": self.project.id, "workflow_id": "wf-art"},
        )
        self.assertEqual(response.status_code, 200)
        stages = {item["stage"]: item for item in response.data["stages"]}
        artifact = stages[PLAN]["output"]["artifact"]
        self.assertTrue(artifact["available"])
        self.assertEqual(artifact["source"], "fallback")
        self.assertTrue(artifact["name"].endswith(".md"))
        # 未产出的阶段不该有这个字段 —— 前端据此**不渲染**按钮（而不是渲染成禁用）。
        self.assertIsNone(stages[WORKFLOW_STAGE_ORDER[1]]["output"])


@override_settings(MEDIA_ROOT=TEST_MEDIA_ROOT)
class StageFeedbackTests(SkillHubBaseTests):
    """「上传反馈」：记录采纳率，**不派生**。"""

    def setUp(self):
        super().setUp()
        from rest_framework.test import APIClient

        self.client = APIClient()
        self.client.force_authenticate(self.executor)
        self._skill, self.version = self.make_skill_version(name="plan-skill", stage=PLAN)
        self.output = self.make_output(
            task_type=PLAN, skill_version=self.version, workflow_id="wf-fb",
        )
        WorkflowGateService.register_output(self.output)

    def _upload(self, data: bytes, *, workflow_id="wf-fb", name="方案报告.xlsx", **extra):
        from django.core.files.uploadedfile import SimpleUploadedFile

        payload = {
            "project": self.project.id, "workflow_id": workflow_id, "stage": PLAN,
            "file": SimpleUploadedFile(name, data, content_type="application/vnd.ms-excel"),
        }
        payload.update(extra)
        return self.client.post(FEEDBACK_URL, payload, format="multipart")

    def test_records_acceptance_rate_bound_to_output_and_version(self):
        response = self._upload(make_stage_report(0.86))

        self.assertEqual(response.status_code, 201)
        self.assertEqual(response.data["acceptance_score"], 86.0)
        self.assertFalse(response.data["below_reference"])
        self.assertTrue(response.data["created"])

        event = FeedbackEvent.objects.get(id=response.data["feedback_id"])
        self.assertEqual(event.reason_code, "report_acceptance_rate")
        self.assertEqual(event.signal, "accepted")
        self.assertAlmostEqual(event.value, 0.86, places=4)
        # 版本溯源靠这两个外键：产出被清理后，反馈仍要能回答"针对哪一版"。
        self.assertEqual(event.output_id, self.output.pk)
        self.assertEqual(str(event.skill_version_id), str(self.version.id))

    def test_low_rate_is_recorded_not_blocked(self):
        """低于参考线照常入库，只标 ``below_reference``。

        采纳率是版本间对比的评分维度，不是准入门槛；低分要可见，但不能拦住记录。
        """
        response = self._upload(make_stage_report(0.42))

        self.assertEqual(response.status_code, 201)
        self.assertEqual(response.data["acceptance_score"], 42.0)
        self.assertTrue(response.data["below_reference"])
        self.assertEqual(response.data["acceptance_reference"], DEFAULT_ACCEPTANCE_REFERENCE)
        self.assertEqual(FeedbackEvent.objects.filter(project=self.project).count(), 1)

    def test_same_report_uploaded_twice_yields_one_feedback(self):
        """同一份报告重复上传不产生第二条反馈。

        幂等键含报告 sha256：重复提交是同一个事件，避免"上传两次变成两条记录"
        把采纳率的版本对比算成两个样本。
        """
        first = self._upload(make_stage_report(0.86))
        second = self._upload(make_stage_report(0.86))

        self.assertEqual(first.status_code, 201)
        self.assertEqual(second.status_code, 201)
        self.assertTrue(first.data["created"])
        self.assertFalse(second.data["created"])
        self.assertEqual(first.data["feedback_id"], second.data["feedback_id"])
        self.assertEqual(FeedbackEvent.objects.filter(project=self.project).count(), 1)

    def test_a_changed_report_is_a_new_event(self):
        """换了文件（哪怕只改一个格）就是一次新评审，不该被去重吃掉。"""
        self._upload(make_stage_report(0.86))
        self._upload(make_stage_report(0.60))

        self.assertEqual(FeedbackEvent.objects.filter(project=self.project).count(), 2)

    def test_report_without_acceptance_label_is_rejected(self):
        """报告末页没有「采纳率」→ 400，并说清是"文件不对"而非"质量不够"。"""
        response = self._upload(make_stage_report(None))

        self.assertEqual(response.status_code, 400)
        self.assertIn("采纳率", detail_text(response))
        self.assertFalse(FeedbackEvent.objects.filter(project=self.project).exists())

    def test_non_excel_upload_is_rejected(self):
        """Word / 任意二进制不是"读不到"，而是**明确不是**这份约定的东西。"""
        response = self._upload(b"not an excel at all", name="report.docx")

        self.assertEqual(response.status_code, 400)
        self.assertIn("Excel", detail_text(response))

    def test_feedback_does_not_derive_or_touch_skill_package(self):
        """反馈只记不派生：Skill 版本数不变、归因一条不落、包哈希不变。

        这条是"反馈与派生解耦"的护栏。合并成一个动作，用户想"只记录一次评审"
        时会被迫改包；而改包是不可逆的。
        """
        baseline_sha = self.version.package_sha256
        baseline_versions = SkillVersion.objects.filter(skill=self._skill).count()

        self._upload(make_stage_report(0.86))

        self.version.refresh_from_db()
        self.assertEqual(self.version.package_sha256, baseline_sha)
        self.assertEqual(SkillVersion.objects.filter(skill=self._skill).count(),
                         baseline_versions)
        self.assertFalse(FailureAttribution.objects.filter(project=self.project).exists())

    def test_stage_without_output_is_rejected(self):
        """没有产出（没有"这一版"）时记不了分，要说清楚而不是落无主反馈。"""
        response = self._upload(make_stage_report(0.86), workflow_id="wf-none")

        self.assertEqual(response.status_code, 400)
        self.assertIn("暂无产出", detail_text(response))
        self.assertFalse(FeedbackEvent.objects.filter(project=self.project).exists())

    def test_missing_file_is_rejected(self):
        response = self.client.post(
            FEEDBACK_URL,
            {"project": self.project.id, "workflow_id": "wf-fb", "stage": PLAN},
            format="multipart",
        )
        self.assertEqual(response.status_code, 400)

    def test_acceptance_history_lists_versions_for_comparison(self):
        """历史采纳率按版本可读 —— 「采纳率作为版本对比维度」的读侧。"""
        self._upload(make_stage_report(0.86))

        history = WorkflowStageFeedbackService.acceptance_history(self._skill)
        self.assertEqual(len(history), 1)
        self.assertEqual(history[0]["version"], self.version.version)
        self.assertAlmostEqual(history[0]["score"], 86.0)
