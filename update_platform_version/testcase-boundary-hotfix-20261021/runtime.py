#!/usr/bin/env python3
"""Runs with the target container's Python; no downloads or pip required."""
import argparse
import ast
import hashlib
import importlib.util
import json
import os
from pathlib import Path
import sys

TARGET = '/app/testcases/review_service.py'
MARKER = '# WHARTTEST_REVIEW_BOUNDARY_HOTFIX_20261021_V1'

def sha(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()

def smoke(candidate):
    os.environ.setdefault('DJANGO_SETTINGS_MODULE', 'wharttest_django.settings')
    sys.path.insert(0, '/app')
    import django
    django.setup()
    from io import BytesIO
    from types import SimpleNamespace
    from openpyxl import load_workbook
    name = 'testcases._review_boundary_preflight'
    spec = importlib.util.spec_from_file_location(name, str(candidate))
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    output = {}
    def save(filename, content, save=False):
        output['bytes'] = content.read()
    review = SimpleNamespace(source_name='preflight.xlsx', skill_name='test-case-clarity-review',
        business_context='', custom_rules='', get_review_mode_display=lambda: '预检',
        report_file=SimpleNamespace(save=save))
    axes = ['正向', '反向', '极值', '异常', '兼容性', '移动响应式', '回归']
    analysis = {'scenarios': [{'module': '登录', 'scenario': '账号登录', 'priority': 'P1',
        'cases': {'正向': ['TC-1']}, 'exemptions': {'移动响应式': '输入明确为非 UI 接口'}}],
        'gaps': [], 'conclusion': '预检'}
    module._write_report(review, [], [], [], [], boundary_analyses=[analysis])
    book = load_workbook(BytesIO(output['bytes']))
    sheet = book['测试场景边界分析']
    assert sheet.sheet_state == 'visible'
    assert sheet.sheet_properties.tabColor.rgb == '00F4B183'
    assert sheet['A1'].fill.fgColor.rgb == '007F6000'
    assert sheet['B8'].value == '=IF(B6=0,0,B7/B6)'
    header = next(row[0].row for row in sheet if row[0].value == '模块')
    assert [sheet.cell(header, n).value for n in range(4, 11)] == axes
    assert sheet.cell(header+1, 4).value == '✓ TC-1'
    assert sheet.cell(header+1, 4).fill.fgColor.rgb == '00E2F0D9'
    assert sheet.cell(header+1, 5).fill.fgColor.rgb == '00FCE4D6'
    assert sheet.cell(header+1, 9).fill.fgColor.rgb == '00E7E6E6'
    assert sheet['B5'].value == 1 and sheet['B6'].value == 6 and sheet['B7'].value == 1
    book.close()
    # Missing model analysis must still produce a visible sheet, without fabricated coverage.
    module._write_report(review, [], [], [], [], boundary_analyses=[])
    book = load_workbook(BytesIO(output['bytes']))
    assert book['测试场景边界分析']['B5'].value == 0
    book.close()
    # Verify the model request actually asks for boundary_analysis, without calling an LLM.
    calls = []
    module.safe_llm_invoke = lambda *args, **kwargs: (calls.append(args[1]) or SimpleNamespace(content='{}'))
    module._review_chunk(object(), 'preflight', [], '')
    assert 'boundary_analysis' in calls[0][1].content
    print('PASS: report sheet, tab/title colors, matrix highlights, formula, empty-result handling, request schema')

def prepare(root, source):
    root = Path(root)
    target = Path(source)
    original = target.read_text(encoding='utf-8')
    patch = json.loads((root/'payload/changes.json').read_text(encoding='utf-8'))
    if MARKER in original:
        print('ALREADY_APPLIED')
        smoke(target)
        return
    if 'workbook.create_sheet("测试场景边界分析")' in original:
        raise RuntimeError('目标已有边界分析导出逻辑；请勿覆盖，请先核验正在执行的 Worker 和报告生成时间。')
    candidate = original
    for index, block in enumerate(patch['blocks'], 1):
        count = candidate.count(block['old'])
        if count != 1:
            raise RuntimeError('目标代码与补丁不匹配：片段 %s 命中 %s 次。未修改业务文件，请提供备份目录中的 original.py 和诊断信息适配。' % (index, count))
        candidate = candidate.replace(block['old'], block['new'], 1)
    candidate += '\n' + MARKER + '\n'
    ast.parse(candidate)
    old_tree, new_tree = ast.parse(original), ast.parse(candidate)
    def functions(tree):
        return {node.name: ast.dump(node, include_attributes=False) for node in tree.body if isinstance(node, (ast.FunctionDef, ast.AsyncFunctionDef))}
    before, after = functions(old_tree), functions(new_tree)
    changed = {key for key in set(before)|set(after) if before.get(key) != after.get(key)}
    assert changed == {'_review_chunk', '_write_report', 'run_testcase_review'}, changed
    (root/'candidate.py').write_text(candidate, encoding='utf-8')
    (root/'metadata.json').write_text(json.dumps({'patch_id':patch['id'], 'target':TARGET,
        'original_sha256':sha(target), 'patched_sha256':sha(root/'candidate.py'),
        'reference_sha256':patch['reference_sha256'], 'changed_functions':sorted(changed)}, indent=2), encoding='utf-8')
    smoke(root/'candidate.py')
    print('PRECHECK_OK')

def host_target(mount_file):
    mounts = json.loads(Path(mount_file).read_text())
    candidates = []
    for mount in mounts:
        destination = mount.get('Destination', '').rstrip('/')
        if TARGET == destination or TARGET.startswith(destination + '/'):
            candidates.append((len(destination), mount, destination))
    if not candidates:
        return
    _, mount, destination = max(candidates, key=lambda row: row[0])
    if mount.get('Type') != 'bind':
        raise RuntimeError('目标源码位于 Docker volume，需按该环境适配；本补丁不覆盖 volume。')
    print(str(Path(mount['Source']) / TARGET[len(destination):].lstrip('/')))

if __name__ == '__main__':
    parser = argparse.ArgumentParser()
    parser.add_argument('action', choices=['prepare','smoke','host-target'])
    parser.add_argument('path')
    parser.add_argument('--source', default=TARGET)
    args = parser.parse_args()
    if args.action == 'prepare': prepare(args.path, args.source)
    elif args.action == 'smoke': smoke(Path(args.path))
    else: host_target(args.path)
