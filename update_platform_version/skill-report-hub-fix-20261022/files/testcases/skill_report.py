"""Declarative extra worksheets; shared renderer, no Skill-specific sheet names."""
import json
import re
from collections import OrderedDict

import yaml
from django.core.exceptions import ValidationError
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

MAX_SHEETS = 12
MAX_COLUMNS = 40
MAX_ROWS = 20000


def _color(value, default):
    value = str(value or '').lstrip('#').upper()
    return value if re.fullmatch(r'[0-9A-F]{6}', value) else default


def _text(value):
    if value is None:
        return ''
    return (json.dumps(value, ensure_ascii=False) if isinstance(value, (dict, list)) else str(value))[:32767]


def _cell(cell, value):
    # Model content is text, including strings that Excel would interpret as formulas.
    cell.value = _text(value)
    cell.data_type = 's'


def validate_sheets(sheets):
    if not isinstance(sheets, list) or len(sheets) > MAX_SHEETS:
        raise ValidationError('report.sheets 必须是数组，最多 12 个工作表')
    normalized, names = [], set()
    for sheet in sheets:
        if not isinstance(sheet, dict):
            raise ValidationError('工作表声明必须是对象')
        name = str(sheet.get('name') or '').strip()
        if not name or len(name) > 31 or re.search(r'[\\/*?:\[\]]', name) or name.startswith("'") or name.endswith("'"):
            raise ValidationError('工作表名称不能为空、不能超过 31 字或包含 Excel 禁用字符')
        if name.casefold() in names:
            raise ValidationError('工作表名称重复：' + name)
        names.add(name.casefold())
        columns = sheet.get('columns')
        if not isinstance(columns, list) or not 1 <= len(columns) <= MAX_COLUMNS:
            raise ValidationError('工作表 %s 必须声明 1–40 列' % name)
        keys, parsed = set(), []
        for column in columns:
            if isinstance(column, str):
                column = {'key': column, 'title': column}
            if not isinstance(column, dict) or not isinstance(column.get('key'), str) or not column['key'].strip():
                raise ValidationError('工作表列必须有非空 key')
            key = column['key'].strip()
            if key in keys:
                raise ValidationError('工作表列 key 重复：' + key)
            keys.add(key)
            try:
                width = max(10, min(100, float(column.get('width', 32))))
            except (TypeError, ValueError):
                raise ValidationError('列宽必须为数字')
            parsed.append({'key': key, 'title': str(column.get('title') or key), 'width': width})
        rules = sheet.get('highlight', [])
        if not isinstance(rules, list) or len(rules) > 50:
            raise ValidationError('highlight 必须是数组，最多 50 条规则')
        for rule in rules:
            if not isinstance(rule, dict) or rule.get('column') not in keys or not ('equals' in rule or 'contains' in rule):
                raise ValidationError('高亮规则必须指定有效 column 和 equals/contains')
        normalized.append({**sheet, 'name': name, 'columns': parsed,
            'tab_color': _color(sheet.get('tab_color'), 'F4B183'),
            'header_color': _color(sheet.get('header_color'), '2F75B5'), 'highlight': rules})
    return normalized


def declared_sheets(skill_prompt):
    match = re.match(r'\s*---\s*\n(.*?)\n---(?:\n|$)', skill_prompt or '', re.S)
    if not match:
        return []
    try:
        manifest = yaml.safe_load(match.group(1)) or {}
    except yaml.YAMLError as exc:
        raise ValidationError('Skill 报告声明 YAML 无法解析') from exc
    report = manifest.get('report', {}) if isinstance(manifest, dict) else {}
    if not isinstance(report, dict):
        raise ValidationError('Skill report 必须为对象')
    return validate_sheets(report.get('sheets', []))


def render_sheets(workbook, declarations, batches):
    definitions = OrderedDict((s['name'], s) for s in declarations)
    rows_by_name, seen, warnings = {}, {}, []
    for batch in batches:
        if not isinstance(batch, list):
            warnings.append('模型的 report_sheets 不是数组，已拒绝该批扩展数据')
            continue
        for item in batch:
            if not isinstance(item, dict):
                warnings.append('模型返回了无效工作表对象')
                continue
            name = str(item.get('name') or '')
            if name not in definitions:
                try:
                    definition = validate_sheets([item])[0]
                except ValidationError as exc:
                    warnings.append('扩展工作表声明无效：' + '; '.join(exc.messages))
                    continue
                if len(definitions) >= MAX_SHEETS:
                    warnings.append('扩展工作表超过 12 个，已拒绝：' + name)
                    continue
                definitions[name] = definition
            elif item.get('columns'):
                try:
                    actual = validate_sheets([item])[0]
                    if [c['key'] for c in actual['columns']] != [c['key'] for c in definitions[name]['columns']]:
                        warnings.append('工作表 %s 的列与已声明结构不一致，已拒绝该批数据' % name)
                        continue
                except ValidationError as exc:
                    warnings.append('工作表 %s 列声明无效：%s' % (name, '; '.join(exc.messages)))
                    continue
            rows = item.get('rows', [])
            if not isinstance(rows, list):
                warnings.append('工作表 %s 的 rows 不是数组' % name)
                continue
            dest, fingerprints = rows_by_name.setdefault(name, []), seen.setdefault(name, set())
            for row in rows:
                if not isinstance(row, dict):
                    warnings.append('工作表 %s 存在非对象数据行，已拒绝' % name)
                    continue
                key = json.dumps(row, ensure_ascii=False, sort_keys=True)
                if key not in fingerprints:
                    if len(dest) >= MAX_ROWS:
                        warnings.append('工作表 %s 超过 20000 行' % name)
                        break
                    dest.append(row)
                    fingerprints.add(key)
    rendered = []
    existing = {name.casefold() for name in workbook.sheetnames}
    for name, definition in definitions.items():
        if name.casefold() in existing:
            warnings.append('扩展工作表与已有工作表重名，未覆盖：' + name)
            continue
        sheet = workbook.create_sheet(name)
        sheet.sheet_properties.tabColor = definition['tab_color']
        sheet.sheet_view.showGridLines = False
        columns = definition['columns']
        sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max(2, len(columns)))
        _cell(sheet.cell(1, 1), definition.get('title') or name)
        sheet.cell(1, 1).font = Font(bold=True, size=16, color='FFFFFF')
        sheet.cell(1, 1).fill = PatternFill('solid', fgColor=definition['header_color'])
        sheet.row_dimensions[1].height = 36
        for n, col in enumerate(columns, 1):
            sheet.column_dimensions[get_column_letter(n)].width = col['width']
            cell = sheet.cell(3, n)
            _cell(cell, col['title'])
            cell.font = Font(bold=True, color='FFFFFF')
            cell.fill = PatternFill('solid', fgColor=definition['header_color'])
            cell.alignment = Alignment(wrap_text=True, vertical='center')
        rows = rows_by_name.get(name, [])
        if not rows:
            _cell(sheet.cell(4, 1), '本次模型未返回该工作表的数据，请复核 Skill 输出；不代表无问题或已覆盖。')
            sheet.cell(4, 1).fill = PatternFill('solid', fgColor='FFF2CC')
            warnings.append('工作表未返回数据：' + name)
        for n, row in enumerate(rows, 4):
            for index, column in enumerate(columns, 1):
                cell = sheet.cell(n, index)
                _cell(cell, row.get(column['key']))
                cell.alignment = Alignment(wrap_text=True, vertical='top')
                cell.fill = PatternFill('solid', fgColor='F7FAFC' if n % 2 == 0 else 'FFFFFF')
            for rule in definition['highlight']:
                value = _text(row.get(rule['column']))
                matched = value == _text(rule['equals']) if 'equals' in rule else _text(rule['contains']) in value
                if matched:
                    targets = list(sheet[n]) if rule.get('scope') == 'row' else [sheet.cell(n, next(i for i,c in enumerate(columns,1) if c['key']==rule['column']))]
                    for cell in targets:
                        cell.fill = PatternFill('solid', fgColor=_color(rule.get('fill'), 'FFF2CC'))
                        cell.font = Font(color=_color(rule.get('color'), '9C6500'), bold=bool(rule.get('bold', True)))
            sheet.row_dimensions[n].height = 54
        sheet.freeze_panes = 'A4'
        sheet.auto_filter.ref = 'A3:%s%d' % (get_column_letter(len(columns)), max(4, sheet.max_row))
        rendered.append(name)
        existing.add(name.casefold())
    return {'extra_sheets': rendered, 'report_warnings': list(dict.fromkeys(warnings))}
